"""Read-only audit of the supplied course template and its native Word render."""
from pathlib import Path
from docx import Document
import json, hashlib, os
from pdf2image import convert_from_path
source = Path(os.environ.get('CHENGVAULT_TEMPLATE', r'D:\Codes\soft-project\课程资料\小组提交文档-实验报告模板.docx'))
output = Path(os.environ.get('CHENGVAULT_TEMPLATE_RENDER', str(Path(__file__).resolve().parents[1] / '.cache' / 'template-render')))
output.mkdir(parents=True, exist_ok=True)
doc = Document(source)
def table_data(table):
    rows=[]
    for i,row in enumerate(table.rows):
        cells=[];seen=set()
        for j,cell in enumerate(row.cells):
            if id(cell._tc) in seen: continue
            seen.add(id(cell._tc))
            cells.append({'column':j,'text':cell.text,'tables':[table_data(t) for t in cell.tables]})
        rows.append({'row':i,'cells':cells})
    return rows
evidence={'source':str(source),'sha256':hashlib.sha256(source.read_bytes()).hexdigest(),
 'paragraphs':[{'index':i,'style':p.style.name,'text':p.text} for i,p in enumerate(doc.paragraphs)],
 'sections':[{'width':s.page_width,'height':s.page_height,'top':s.top_margin,'bottom':s.bottom_margin,'left':s.left_margin,'right':s.right_margin,'header':s.header_distance,'footer':s.footer_distance} for s in doc.sections],
 'tables':[table_data(t) for t in doc.tables]}
(output/'structure.json').write_text(json.dumps(evidence,ensure_ascii=False,indent=2),encoding='utf-8')
pdf=convert_from_path(str(output/'template.pdf'),dpi=96,poppler_path=os.environ.get('POPPLER_PATH', r'C:\Program Files\poppler\Library\bin'))
for index,page in enumerate(pdf):
    page.save(output/f'page-{index+1:02}.png')
print('Pages:',len(pdf),'Tables:',len(doc.tables),'Paragraphs:',len(doc.paragraphs))
for i,p in enumerate(doc.paragraphs):
    if p.text: print('P',i,repr(p.text[:150]))
for i,row in enumerate(doc.tables[0].rows):
    print('ROW',i,[(j,cell.text[:160].replace('\n',' / '),len(cell.tables)) for j,cell in enumerate(row.cells)])
