"""Build the course report in the retained DOCX template using measured evidence.

Only designated table slots and new image relationships change. Preserve-only
ZIP parts are restored byte-for-byte from the reference after python-docx edits.
Run with the bundled Python runtime; never modifies the original template.
"""
from __future__ import annotations
import argparse, csv, hashlib, json, re, subprocess, zipfile
from collections import Counter
from datetime import datetime, timezone, timedelta
from pathlib import Path
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
REFERENCE = ROOT.parent / '文档/小组提交文档-实验报告模板.docx'
OUTPUT = ROOT.parent / '报告文档'
EXPECTED_REFERENCE = '3323e15ccfa3445483575fd6b77a1606c8b7c64cbdc30cca69cdea858913a2e0'
CHINA = timezone(timedelta(hours=8))

REQS = [
('B01','目录树备份','保留普通文件、空文件、空目录、中文路径及层次结构','filesystem.cpp','fs.roundtrip'),
('B02','指定位置恢复','恢复字节一致，默认不覆盖，失败保护已有文件','filesystem.cpp','fs.restore'),
('E01','NTFS 特殊文件','符号链接、硬链接、Junction、稀疏文件、文件及目录 ADS','filesystem.cpp','fs.ntfs'),
('E02','NTFS 元数据','FILETIME、属主和组 SID、DACL 及继承、普通属性','filesystem.cpp','fs.metadata'),
('E03','路径筛选','包含和排除相对路径；排除优先','filesystem.cpp','fs.filter.path'),
('E04','类型筛选','对象类型与扩展名组合','filesystem.cpp','fs.filter.object'),
('E05','名字筛选','名字通配符与祖先目录保留','filesystem.cpp','fs.filter.name'),
('E06','时间筛选','创建 访问 修改时间的包含式区间','filesystem.cpp','fs.filter.time'),
('E07','尺寸筛选','逻辑字节数上下界均包含边界','filesystem.cpp','fs.filter.size'),
('E08','用户筛选','文件属主 SID 或 Windows 账户名','filesystem.cpp','fs.filter.owner'),
('E09','额外筛选','属性、深度、UTF8 内容、相对快照变化','filesystem.cpp','fs.filter'),
('E10','索引打包解包','内容块索引、去重、独立单文件完整恢复','archive.cpp','archive.roundtrip.indexed'),
('E11','PAX 打包解包','真实 TAR 成员与 NTFS 扩展，独立读取器验证','archive.cpp','archive.pax'),
('E12','Huffman 压缩解压','自主前缀编码，频数与边界校验','compression.cpp','Huffman/'),
('E13','LZ77 压缩解压','自主滑动窗口匹配与重叠回溯','compression.cpp','LZ77/'),
('E14','AES256 加密解密','自主 AES256 CTR 与 HMAC 认证，用户指定密码','crypto.cpp','AES256CTR/'),
('E15','ChaCha20 加密解密','自主 ChaCha20 与 HMAC 认证，拒绝篡改','crypto.cpp','ChaCha20/'),
('E16','GUI','八页中文界面，异步耗时操作、进度及错误信息','main_window.cpp','GUI/'),
('E17','定时及淘汰','间隔 每日 每周任务、漏跑补偿、版本保留','scheduler.cpp','SCH-'),
('E18','实时备份','Windows 事件、合并去抖、增量提交','scheduler.cpp','SCH-14'),
('E19','网盘式备份','独立服务端、上传、列举、下载并恢复','network.cpp','network.full'),
('E20','用户管理','创建 禁用 密码重置 配额 设备授权与撤销','network.cpp','network.admin'),
('E21','远程元数据','用户隔离索引；文件级清单保持加密','network.cpp','network.account'),
('E22','加密传输','设备预共享密钥认证、方向密钥、序号防重放','network.cpp','network.replayed'),
('E23','增量及续传','只发送缺失密文对象；中断按偏移续传','network.cpp','network.changed'),
('E24','冗余纠错扩展','自主 RS12 10，单组至多两个已定位分片损坏可修复','repair.cpp','RS12-10/'),
('Q01','健壮性和安全','越界、路径穿越、错误密码、取消、占用均明确失败','filesystem.cpp','fs.manifest'),
('Q02','可构建与可维护','CMake Ninja 自动构建，模块化、命名及原创注释','CMakeLists.txt','Build/'),
('Q03','真实验收证据','白盒测试及至少八类测试，证据可重跑','test_main.cpp',''),
('D01','课程完整交付','源码 程序 构建脚本 模板报告 PPT 视频 UML','docs',''),
]

def load_json(path: Path, default=None):
    if not path.exists(): return {} if default is None else default
    return json.loads(path.read_text(encoding='utf-8-sig'))

def choose_evidence():
    evidence = load_json(ROOT / 'evidence/test-results.json')
    cases = evidence.get('cases', evidence.get('tests', []))
    assert cases, 'No executed test evidence available'
    return evidence, cases

# These descriptions are reviewed against the corresponding C++ assertions.
# They are explanatory metadata; only executed JSON determines PASS/FAIL/SKIP.
EXTRA_DESCRIPTIONS = {}
for description_file in (ROOT/'evidence').glob('report-*-case-descriptions.json'):
    EXTRA_DESCRIPTIONS.update(load_json(description_file))

def describe(name: str):
    """Human-readable inputs/actions/assertions grounded in the actual tests."""
    if name in EXTRA_DESCRIPTIONS: return EXTRA_DESCRIPTIONS[name]
    fixed = {
    'Common/bounded-read-exact-and-over-limit':('读取长度边界','写入 7 字节 bounded，分别以 7、6、-1、int64 最大值读取。','仅上限 7 的读取成功且字节一致；其余明确拒绝。'),
    'Common/empty-file-zero-bound':('空文件读取','零字节文件以零上限读取，再读取缺失路径。','空文件返回空字节；缺失路径报错。'),
    'Common/json-object-required':('JSON 对象约束','依次写入 []、null、{，再写 accepted=true 对象。','前三种拒绝；合法对象字段保持。'),
    'Credentials/DPAPI-current-user-roundtrip':('当前用户凭据','随机 32 字节密钥经当前用户 DPAPI 封装解封，再截短封装到 10 字节。','封装不同于明文且完整往返；截短封装拒绝。'),
    'Crypto/invalid-key-nonce':('密码参数边界','AES 短密钥、ChaCha 短 nonce、seal 不支持 none 及短主密钥；比较不同长度标签。','非法参数拒绝；相同标签 true，不同内容或长度 false。'),
    'Entropy/OS-random-length-and-distinct':('系统随机数','请求 0 字节、两次 32 字节、负 1 字节。','零长为空；两次随机输出不同；负长度拒绝。'),
    'HashFile/stream-and-cancel':('文件流式摘要','实际文件 hashFile 与原始字节 SHA256 比较，再以取消标志执行。','摘要一致；预先取消时抛错。'),
    'KDF/invalid-parameters':('密钥派生边界','PBKDF2 迭代 0 或输出 0；HKDF 输出 8161 或 0 字节。','非法迭代及输出长度全部拒绝。'),
    'fs.roundtrip.files-empty-unicode':('普通文件 空目录 中文名称','建立 hello.txt、空文件、空目录和中文路径的二进制文件；捕获后恢复到新目录。','文件数量为 3，目录存在，文本与二进制逐字节一致。'),
    'fs.metadata.owner-group-times-dacl':('属主 权限 时间','捕获源文件，再恢复并重新捕获目标元数据；逐字段比较。','ownerSid、groupSid、创建与修改时间、属性和 DACL 一致。'),
    'fs.filter.path-exclusion-wins':('路径包含与排除','包含 source/*，排除 source/nested，再预览。','仅根目录的两个文件命中，nested 子树不被选中。'),
    'fs.filter.name-glob-ancestor':('名字通配符','名字设为 中文.*，捕获并检查祖先。','只选择中文文件，同时保留 source 和 nested 祖先目录。'),
    'fs.filter.object-type-and-extension':('类型与扩展名','类型设 file，扩展名设 .TXT，执行预览。','大小写不敏感地命中两个 txt 文件。'),
    'fs.filter.size-inclusive':('尺寸精确边界','最小和最大尺寸都设为 7 字节。','仅 7 字节中文文件命中。'),
    'fs.filter.owner-sid':('属主 SID','先用实际 owner SID 筛选，再改成 S-1-0-0。','实际 SID 选择全部三个文件，错误 SID 不选择文件。'),
    'fs.filter.owner-account-name':('Windows 属主账户名','读出 DOMAIN\\user 形式账户名并用它筛选。','反斜杠不被误当路径转换；三个文件均命中。'),
    'fs.filter.time-bounds':('时间区间','下界设明天，再将上界设明天，分别预览。','前者无命中；后者三个文件命中。'),
    'fs.filter.invalid-time-rejected':('非法时间','将 timeFrom 设 not-a-date。','拒绝配置，不产生成功预览。'),
    'fs.filter.depth':('递归深度','maxDepth 设为 1，根目录深度为 0。','只选根目录中的两个文件，不选深层中文文件。'),
    'fs.filter.attributes':('隐藏属性','给 hello.txt 设置 HIDDEN，再要求必须含 HIDDEN 位。','仅隐藏文件命中。'),
    'fs.filter.content-boundary':('跨块内容匹配','将 NEEDLE 放在 1 MiB 读取边界两侧并筛选。','重叠窗口找到完整字符串，不漏选边界匹配。'),
    'fs.filter.changed-since':('相对快照变化','用旧清单筛选，再改写 hello.txt 内容与大小。','未变化时无命中；改写后只命中该文件。'),
    'fs.restore.no-overwrite-preflight':('默认保护已有文件','目标预置 keep original，再恢复且不启用覆盖。','任务失败；原字节保留，其他目标子目录尚未创建。'),
    'fs.restore.explicit-overwrite':('显式覆盖','先恢复，再改写目标为 old，启用覆盖后再恢复。','同目录临时文件提交后，内容恢复为源字节。'),
    'fs.restore.disk-full-write-boundary-injection':('写入中途故障注入','写入首块后通过进度回调抛 DISK_FULL_SIMULATED。','任务失败、临时文件清理，旧目标保持；此项不是实际填满磁盘。'),
    'fs.capture.denied-dacl-restricted-token':('真实拒绝访问 DACL','给临时文件加入拒绝 READ_DATA 的 DACL，以受限令牌捕获。','读取拒绝导致失败；结束后恢复原 DACL 与令牌。'),
    'fs.restore.selection':('选择目录恢复','仅选择 source/nested 进行恢复。','中文文件存在，未选中的 hello.txt 不出现。'),
    'fs.restore.corruption-no-destination':('损坏数据块','捕获后翻转块中字节，调用恢复。','预检失败且恢复根目录未创建。'),
    'fs.capture.cancelled':('预先取消','取消标志为 true 后执行 capture。','检查点抛取消错误，不返回成功清单。'),
    'fs.capture.source-target-overlap':('源目标重叠','将归档目标放在源目录内部。','拒绝备份，避免归档递归包含自身。'),
    'fs.capture.exclusively-locked-source':('独占文件占用','用零共享模式持有源文件句柄，再捕获。','明确失败，不静默跳过被占用文件。'),
    'fs.manifest.case-collision':('大小写冲突','向清单加入已有 hello.txt 的 HELLO.TXT 副本。','验证阶段拒绝冲突。'),
    'fs.ntfs.hardlinks-shared-identity':('硬链接身份','为 hello.txt 建立 hard.txt，恢复后读取文件 ID。','两个路径共享 NTFS ID；旧快照筛选不重复误判变化。'),
    'fs.ntfs.file-and-directory-ads':('文件与目录 ADS','分别写入文件流和目录 origin 流，捕获后恢复。','两种命名流的名称与字节均一致。'),
    'fs.ntfs.sparse-over-4gib':('超过 4 GiB 稀疏文件','建立 2^32+123 字节稀疏文件，只写末尾 TAIL。','存储数据小于 1 MiB；恢复长度、SPARSE 标志和尾部字节正确。'),
    'fs.ntfs.junction-no-recursion':('目录联接','Junction 指向 nested；捕获、恢复并读取属性。','清单无联接下重复子树；恢复后仍是重解析点。'),
    'fs.restore.reparse-escape-rejected':('重解析点逃逸','使恢复根目录成为指向 outside 的 Junction。','恢复失败且 outside 中没有新建备份树。'),
    'fs.ntfs.symbolic-file-and-directory-links':('符号链接','建立相对文件与目录符号链接，备份后恢复。','链接保持 REPARSE_POINT，文件链接读取与原始目标一致；无权限则跳过。'),
    'archive.wrong-password-no-output':('归档错误密码','用正确密码生成 ChaCha20 归档，再用错误密码恢复。','认证失败且不创建恢复目录。'),
    'archive.pax-independent-reader':('PAX 独立读取器','生成 PAX 后解开通用外层，以 Windows tar -tf 列举内部流。','系统 tar 返回 0，能看到 hello.txt 和空目录。'),
    'archive.tampered-final-record':('归档末尾篡改','修改归档最后一字节，再执行 verify。','校验或认证失败。'),
    'archive.truncation-detected':('归档截断','删除归档末尾 80 字节后恢复。','缺少有效终结记录，拒绝恢复且无目标目录。'),
    'repository.initialize.random-master':('仓库主密钥','初始化加密仓库并读取标识与主密钥长度。','生成有效仓库 ID 和独立随机主密钥。'),
    'repository.first-snapshot':('首次快照','捕获初始源文件并查看快照清单。','提交完整快照并报告新增数据块。'),
    'repository.unchanged-blocks-reused':('未变化数据复用','同一源连续备份两次。','第二次不重复保存原始数据块。'),
    'repository.local-modification-one-chunk':('单块修改','只修改一个 1 MiB 区间后再次备份。','只新增变化块，其余块复用。'),
    'repository.snapshot-full-restore':('独立版本恢复','选择完整快照恢复到空目录。','不依赖逐份增量文件回放，内容一致。'),
    'repository.manifest-encrypted':('清单保密','检查仓库中的快照密文文件。','不能直接读取文件名和明文清单内容。'),
    'repository.wrong-password':('仓库错误密码','用错误密码打开已初始化仓库。','解封主密钥失败。'),
    'repository.dpapi-master-unlock':('后台凭据解封','当前用户 DPAPI 保护及解封主密钥后打开仓库。','密钥可用于同一仓库且不需要保存明文口令。'),
    'repository.password-change-rewrap-only':('仓库换密码','记录密文对象，修改口令，再比较并重新打开。','只更新主密钥封装；已有数据块密文不变。'),
    'repository.verify-all-snapshots':('仓库完整性','逐一验证所有活跃快照及其唯一块。','清单和数据块全部通过完整性验证。'),
    'repository.concurrent-task-lock':('并发写锁','持有仓库锁时发起另一个任务。','拒绝冲突写入，避免两个清单交错提交。'),
    'repository.prune-shared-blocks':('版本淘汰','产生多个共享数据的快照后按保留数淘汰。','仅删除无引用对象，剩余快照仍可恢复。'),
    'repository.path-traversal-rejected':('仓库路径越界','请求非法快照或对象相对路径。','路径校验拒绝越界访问。'),
    'repository.cancel-no-new-snapshot':('仓库取消事务','备份过程中设置取消。','不发布新的成功快照。'),
    'repository.tampered-object-rejected':('仓库密文损坏','修改一个持久化对象后读取或验证。','认证失败，不返回伪造明文。'),
    'network.setup-two-users':('真实服务端账户','本机启动 TCP 服务端，配置 Alice 与 Bob 两个账户。','两个独立账户及设备可用于随后认证测试。'),
    'network.first-upload-real-tcp':('首次网络上传','认证 Alice 并上传完整仓库。','实际 TCP 上传对象成功且产生非零数据负载。'),
    'network.unchanged-zero-payload':('重复上传','对已提交仓库原样再次上传。','payloadBytes 为 0；控制帧流量可不为 0。'),
    'network.account-isolation-list':('列表用户隔离','用 Bob 列举 Alice 已上传后的仓库。','Bob 的列表不泄漏 Alice 仓库。'),
    'network.account-isolation-download':('下载用户隔离','Bob 使用 Alice 仓库 ID 下载。','拒绝跨用户读取。'),
    'network.full-download-restore':('网盘下载恢复','Alice 下载仓库、解锁、选择版本恢复。','恢复内容与本地源一致。'),
    'network.download-no-overwrite':('下载目标保护','向非空下载目标发起完整仓库下载。','拒绝覆盖已有本地内容。'),
    'network.changed-upload-incremental':('网络变化块','本地修改一个块、提交快照并重新上传。','复用未变对象，传输量小于完整重传。'),
    'network.wrong-login-password':('登录密码错误','设备配对正确但登录口令错误。','认证业务失败，不获得仓库数据。'),
    'network.wrong-device-key':('设备密钥错误','账户正确但使用错误设备 PSK。','握手认证失败。'),
    'network.device-bound-to-user':('设备账户绑定','尝试把一名用户设备用于另一用户。','绑定校验拒绝认证。'),
    'network.nonadmin-management-denied':('普通用户管理权限','普通账户调用管理员操作。','服务端拒绝越权，不依赖 GUI 隐藏按钮。'),
    'network.admin-lists-no-secrets':('管理列表脱敏','管理员列举账户和设备。','返回管理字段，不返回密码散列或配对密钥。'),
    'network.disable-user':('禁用账户','管理员禁用账户后，该账户继续请求。','后续业务请求失败。'),
    'network.reset-login-password':('重置登录密码','管理员重置口令后用旧口令重新认证。','旧口令失效，仓库内容密码不随之改变。'),
    'network.quota-enforced':('账户配额','将账户配额调小后尝试超额上传。','提交被拒绝，不把未完成上传标记成功。'),
    'network.admin-create-register-device':('动态管理','管理员新建账户并注册设备，再按授权登录。','新账户设备可用，授权在服务端生效。'),
    'network.interrupted-upload-resumes-offset':('真实中断续传','实际取消上传，再对同一对象恢复传输。','确认的暂存偏移被复用，完成后摘要一致。'),
    'network.replayed-frame-rejected':('消息重放','向认证会话重复发送先前有效帧。','严格序号校验拒绝重放。'),
    'network.ciphertext-tamper-rejected':('网络密文篡改','翻转认证密文帧的字节。','HMAC 认证失败，业务消息不被执行。'),
    'network.oversized-frame-rejected':('超长报文','发送超过协议上限的长度字段。','分配内存前拒绝。'),
    'network.server-stores-no-master-or-password':('服务端密钥边界','检查服务端持久化账户与仓库文件。','不存储明文登录密码或备份主密钥。'),
    'network.reject-user-path-traversal':('账户路径注入','使用含目录穿越的账户或请求标识。','请求被拒绝，不访问服务根外的文件。'),
    }
    if name in fixed: return fixed[name]
    if name.startswith('fs.manifest.reject-'):
        bad = name.removeprefix('fs.manifest.reject-')
        return ('恶意路径 '+bad,'将清单路径替换为 '+bad+'，再调用 validateManifest。','在磁盘写入前拒绝设备别名、越界或非法名称。')
    if name.startswith('archive.roundtrip.'):
        _,_,fmt,comp,cipher=name.split('.')
        return (f'{fmt} {comp} {cipher}',f'用 {fmt} 打包、{comp} 压缩及 {cipher} 加密，执行 inspect、verify、restore。','完整目录树和文件字节一致，读取器识别正确内部格式。')
    if name.startswith('SCH-'):
        n=int(name[4:6]); data={
        1:('间隔时钟','注入 2026-10-17 12:30+08:00，间隔 60 秒。','下一次为参考时间加 60 秒。'),
        2:('零间隔','intervalSeconds=0。','钳制为 1 秒，避免立即循环。'),
        3:('负间隔','intervalSeconds=-60。','下一次仍晚于参考时间。'),
        4:('每日未到时刻','参考 12:30，配置 13:00。','保持同一天且时刻为 13:00。'),
        5:('每日已过时刻','参考 12:30，配置 09:00。','安排到下一天。'),
        6:('每日精确边界','参考和计划均为 12:30。','安排下一天而非立即循环。'),
        7:('每周星期','设星期一 09:00。','返回未来的星期一。'),
        8:('非法星期','设 weekday=8。','拒绝无效任务。'),
        9:('非法时刻','设 time=25:00。','拒绝无效任务。'),
        10:('未知模式','mode=invalid。','拒绝配置而非猜测默认。'),
        11:('任务口令保密','设置 plan.password 后调用默认 toJson。','序列化结果不含 password 字段。'),
        12:('任务筛选持久化','多源及深度、内容条件序列化再读取。','源列表和 filters 完全保留。'),
        13:('取消检查点','将原子取消标志置 true 并 check。','立即抛取消异常。'),
        14:('真实写入事件','创建 DirectoryWatch，真实写入 change.txt，轮询事件。','两秒内观察到 ReadDirectoryChangesW 通知。'),
        15:('监听退出','启动监听后销毁对象。','取消挂起 I/O，退出耗时小于 3 秒。'),
        16:('配置读写','原子写有效中文 JSON，再读取非法 JSON。','有效数据保持，损坏配置抛错。')}
        data.update({
        17:('任务密钥保护','保存任务并检查文件字节，再以当前用户 DPAPI 解封。','不含明文口令或主密钥，解封后主密钥一致。'),
        18:('无 GUI 后台执行','启动 Engine，等待持久化逾期任务完成。','生成一个快照并回到 waiting 状态。'),
        19:('漏跑补偿与淘汰','将计划改为过去时间，连续三次修改并运行，保留数为 2。','剩余两个版本且仓库可完整校验。'),
        20:('禁用任务','禁用任务后启动和停止 Engine。','没有额外快照。'),
        21:('实时版本提交','启用 realtime，真实改写文件后处理事件。','产生新完整版本并通过仓库校验。'),
        22:('失败后有界重试','任务源改为不存在目录并运行。','状态变为 retrying，失败计数和错误信息持久化。'),
        23:('删除任务持久化','删除已保存任务并重新列举。','任务列表为空。'),
        24:('任务路径穿越','尝试删除 ../outside 任务 ID。','拒绝越界标识。')})
        data.update({
        25:('损坏任务可见','写入 {broken 作为任务 JSON，再列举和删除。','显示 invalid 且禁用；保留文件名对应 ID，可明确删除。'),
        26:('持久化任务 ID 冲突','wrong-name.json 内写 id=other，再执行 tick。','任务无效，不生成或覆盖 other.json。'),
        27:('非法任务 ID','分别删除空 ID、job:stream、129 字符及 dir\\job。','拒绝空值、ADS、过长与路径分隔符。'),
        28:('不完整计划','分别清空源列表与目标目录后保存任务。','均拒绝，持久任务列表仍为空。'),
        29:('时间与星期下界','传无效参考时间、weekday=0；检查 realtime 与本周未来时刻。','非法值拒绝；实时下一秒，同日未来计划不跳周。'),
        30:('未来任务不提前','nextRun 设参考时间次日，再 tick。','快照数不变，任务保持 waiting。'),
        31:('损坏执行时间补跑','nextRun 设非法文本，连续两次 tick。','只接纳一次运行，完成记录 lastRun，返回 waiting。'),
        32:('坏任务故障隔离','前一任务 mode=unsupported，后一有效任务到期。','有效任务仍完成，坏任务日志含 ID，不阻塞队列。'),
        33:('DPAPI 密钥损坏','持久 key 设 AAAA 后运行。','进入 retrying，failures=1，快照数不增加。'),
        34:('重试上溢边界','failures=2147483647 且密钥无效后执行。','计数饱和到 60，重试约 300 秒，不溢出。'),
        35:('运行中删除任务','接纳任务后立即删除再停止引擎。','完成回调不重新创建已删除 JSON。'),
        36:('运行中配置损坏','接纳后将任务 JSON 改为 {broken，再等完成回调。','损坏文件保持原字节，日志记录任务，不覆盖为成功。'),
        37:('暂停实时监听','已有实时任务连续 tick 后禁用并再次 tick。','释放监听器，未产生额外快照。'),
        38:('单文件监听父目录','监听 selected.txt 后真实改写，随后消费变化标志。','观察到通知；消费完成后再次读取为 false。'),
        39:('监听目录缺失','监听不存在的 missing/child/file。','请求失败重扫，而不是静默无变化。'),
        40:('重命名与删除事件','实际 before.txt 改名为 after.txt，再删除。','两次操作分别触发 ReadDirectoryChangesW 重扫信号。'),
        41:('无效自启动程序','记录注册表启用状态后设置缺失 Agent 路径。','设置失败，注册表原状态不变。'),
        42:('配置目录不可访问','以已有普通文件作为 CHENGVAULT_DATA_HOME，再 tick。','异常不逃出定时回调，原文件字节保持。'),
        43:('完成时失去配置目录','任务运行后切换配置目录为普通文件，执行完成回调。','Agent 不崩溃；原任务仍为 running，不能伪记成功。'),
        44:('重试负数边界','failures=-2147483647 且密钥无效后运行。','计数归一到 1，nextRun 晚于当前时间，不形成忙循环。')})
        return data.get(n,(name,'执行对应调度状态与实际任务操作。','行为符合该测试的具体断言。'))
    if name.startswith('SHA256/'):
        tail=name.split('/',1)[1]
        if 'split-' in tail: return ('SHA256 分块边界 '+tail.split('-')[-1],'按指定长度分次喂入相同原始消息。','增量摘要等于一次性 SHA256，跨块状态正确。')
        return ('SHA256 '+tail,'运行源码中固定 FIPS 消息及标准摘要断言；million 用流式喂入。','计算值与独立标准答案一致；非法 finish 后 update 被拒绝。')
    if name.startswith(('HMAC/','PBKDF2/','HKDF/','AES256CTR/','ChaCha20/')):
        if 'exhaustion' in name: return ('计数器耗尽 '+name.split('/')[0],'将计数器置于极限并请求越界长度。','拒绝溢出，避免复用密码流。')
        return (name,'使用测试源码所列 RFC 或 NIST 固定密钥、消息、盐和期望值计算。','输出逐字节等于标准向量，不只验证自加密自解密往返。')
    if name.startswith('Compression/') and '/fixture-' in name:
        _,alg,fixture=name.split('/')
        return (alg+' 压缩样本 '+fixture.split('-')[-1],'对固定样本集合中的对应字节串 encode，再以原长上限 decode。','还原字节完全一致；样本含空输入、重复、随机及边界数据。')
    if name.startswith('Compression/'):
        return (name.replace('Compression/','压缩 '),'对原始长度、压缩体或算法标识执行测试指定的截断、拼接、伪造或上限修改。','在明确大小上限内解析并拒绝非法输入。')
    if name.startswith(('Huffman/','LZ77/')):
        return (name,'构造频数表损坏、零距离回溯或跨窗口重叠样本，调用编码/解码。','合法边界可逆；无效编码拒绝且不越界。')
    if name.startswith('Envelope/'):
        return ('认证信封 '+name.split('/')[1], '执行指定长度往返，或按测试名修改密钥、头部、密文、标签及尾部。','合法数据原样还原；错误密钥与篡改在返回明文前被拒绝。')
    if name.startswith('RS12-10/single-erasure-'):
        i=name.rsplit('-',1)[1]; return ('RS 单片缺失 '+i,'对系统化编码结果删除第 '+i+' 片并解码。','原始十个数据片逐字节恢复。')
    if name.startswith('RS12-10/double-erasure-'):
        a,b=name.rsplit('-',2)[-2:]; return (f'RS 双片缺失 {a} {b}',f'删除第 {a} 与第 {b} 片，其余分片有效。','有限域逆矩阵重建，全部原始数据片一致。')
    if name.startswith('RS12-10/'):
        return (name.replace('RS12-10/','RS '),'对系统化编码、超过两片缺失、不一致幸存片或非法分片长度执行断言。','合法十片保持原数据；超能力和结构错误拒绝。')
    if name.startswith('Repair/'):
        item=name.split('/',1)[1]
        descriptions={
            'create-real-sidecar':('真实纠错文件','对实际文件生成 .cvr 并检查配置和校验数据。','生成可验证的 RS12 10 侧车。'),
            'two-damaged-data-shards':('两个数据片损坏','翻转同组两个数据片后修复。','新输出与原归档完全一致。'),
            'data-and-parity-damaged':('数据与校验片损坏','损坏同组一个数据片和一个校验片。','利用剩余十片成功恢复。'),
            'both-parity-damaged':('两个校验片损坏','两个校验片无效但十个数据片有效。','原始归档仍可重建并验证。'),
            'three-damaged-reject-no-output':('超过纠错能力','同组损坏三个分片后修复。','明确失败且不发布输出。'),
            'truncated-archive-tail':('归档尾部截断','截掉归档尾部后用有效侧车恢复。','在纠错能力内补回末尾并保持原逻辑长度。'),
            'head-metadata-damaged-tail-used':('首元数据副本损坏','破坏侧车头部元数据。','尾部有效副本提供恢复参数。'),
            'tail-metadata-damaged-head-used':('尾元数据副本损坏','破坏侧车尾部元数据。','头部有效副本提供恢复参数。'),
            'both-metadata-copies-damaged':('双元数据损坏','破坏头尾两份索引。','拒绝猜测编码参数。'),
            'valid-but-conflicting-metadata-copies':('有效副本冲突','构造摘要有效但内容冲突的两份元数据。','拒绝选择性信任冲突副本。'),
            'existing-output-protected':('修复目标保护','预置已有输出文件后调用修复。','原文件不被覆盖。'),
            'cancel-without-output':('修复取消','取消标志生效后开始或继续修复。','不发布不完整结果。'),
            'wrong-sidecar-rejected':('错误配对侧车','使用另一归档的 .cvr。','完整摘要校验失败，拒绝输出。'),
            'empty-archive':('空归档纠错','对空文件生成侧车并修复。','恢复长度为零，结构合法。'),
            'encrypted-archive-still-needs-password':('修复加密数据','对认证密文做纠错恢复，再尝试错误和正确密钥。','纠错不绕过认证；正确密钥恢复原文。'),
        }
        return descriptions.get(item,('纠错 '+item,'按测试源码制造尾部拼接或校验分片截断后执行检查。','损坏被检出，不能产生未经完整摘要验证的成功结果。'))
    return (name,'运行源码中对应白盒/边界输入，检查返回值及异常分支。','断言成立；真实结果以证据中记录为准。')

def requirement_for(name):
    if name.startswith('fs.ntfs'): return 'E01'
    if name.startswith('fs.metadata'): return 'E02'
    if name.startswith('fs.filter'):
        if name.startswith(('fs.filter.glob.literal-directory','fs.filter.glob.star-crosses')): return 'E03'
        if name.startswith(('fs.filter.glob','fs.filter.string-name')): return 'E05'
        if name.startswith('fs.filter.hardlink'): return 'E04'
        for k,v in [('path','E03'),('object','E04'),('name','E05'),('time','E06'),('size','E07'),('owner','E08')]:
            if name.startswith('fs.filter.'+k): return v
        return 'E09'
    if name.startswith('fs.roundtrip'): return 'B01'
    if name.startswith('fs.restore'): return 'B02'
    if name.startswith('fs.'): return 'Q01'
    if name.startswith('archive'): return 'E11' if 'tar' in name or 'pax' in name else 'E10'
    if name.startswith(('AES','Envelope/aes','Envelope/AES')): return 'E14'
    if name.startswith(('ChaCha','Envelope/chacha')): return 'E15'
    if name.startswith(('Compression/huffman','Huffman')): return 'E12'
    if name.startswith(('Compression/lz77','LZ77')): return 'E13'
    if name.startswith(('RS12','Repair')): return 'E24'
    if name.startswith('repository'): return 'E23'
    if name.startswith('network'):
        if name.startswith(('network.account-isolation-list','network.server-stores-no-master')): return 'E21'
        if any(k in name for k in ('account-disable','device-key-rotation')): return 'E20'
        if any(k in name for k in ('replay','tamper','frame','key')): return 'E22'
        if any(k in name for k in ('admin','user','account','password','quota','device')): return 'E20'
        if any(k in name for k in ('incremental','unchanged','resume')): return 'E23'
        return 'E19'
    if name.startswith('SCH'): return 'E18' if name.startswith(('SCH-14','SCH-15','SCH-21','SCH-37','SCH-38','SCH-39','SCH-40')) else 'E17'
    return 'Q03'

def source_for(name, fallback):
    for prefixes,source in [
        (('fs.',),'src/platform/filesystem.cpp'),
        (('archive.',),'src/core/archive.cpp'),
        (('repository.',),'src/core/repository.cpp'),
        (('network.',),'src/core/network.cpp'),
        (('SCH-',),'src/core/scheduler.cpp'),
        (('Common/',),'src/core/common.cpp'),
        (('Compression/','Huffman/','LZ77/'),'src/core/compression.cpp'),
        (('RS12-10/','Repair/'),'src/core/repair.cpp'),
        (('SHA256/','HMAC/','PBKDF2/','HKDF/','AES','ChaCha','Envelope/','KDF/','Crypto/','Entropy/','Credentials/','HashFile/'),'src/core/crypto.cpp')]:
        if name.startswith(prefixes): return source
    return fallback

def style_run(run, size=11, bold=False, font='宋体'):
    run.font.name=font; run.font.size=Pt(size); run.font.bold=bold; run.font.color.rgb=RGBColor(0,0,0)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'),font)
    return run

def para(cell,text='',size=11,bold=False,center=False,indent=True,after=4,keep=False):
    p=cell.add_paragraph(); p.paragraph_format.line_spacing=1.4
    p.paragraph_format.space_after=Pt(after); p.paragraph_format.space_before=Pt(0)
    # Word treats keepNext in this long outer-table template as a request to
    # keep whole adjacent table rows together. Explicitly disable that inherited
    # behavior; headings/images are checked after native Word pagination.
    p.paragraph_format.keep_with_next=False
    p.paragraph_format.keep_together=True
    if indent and not center: p.paragraph_format.first_line_indent=Pt(size*2)
    if center: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    style_run(p.add_run(text),size,bold)
    return p

def heading(cell,text,level=1):
    p=para(cell,text,size=18 if level==1 else 13 if level==2 else 11.5,bold=True,center=level==1,indent=False,after=7,keep=True)
    p.paragraph_format.space_before=Pt(10 if level==1 else 7)
    p.paragraph_format.keep_with_next=(level==1)
    # A real outline level makes navigation work without changing template styles.
    out=OxmlElement('w:outlineLvl'); out.set(qn('w:val'),str(max(0,level-1))); p._p.get_or_add_pPr().append(out)
    return p

def clear_cell(cell):
    for child in list(cell._tc):
        if child.tag != qn('w:tcPr'): cell._tc.remove(child)

def table(cell,headers,rows,widths=None,font_size=10):
    widths=widths or [16.0/len(headers)]*len(headers)
    t=cell.add_table(rows=1, cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for col,w in zip(t.columns,widths): col.width=Cm(w)
    props=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        x=OxmlElement('w:'+edge); x.set(qn('w:val'),'single'); x.set(qn('w:sz'),'4'); x.set(qn('w:color'),'000000'); borders.append(x)
    props.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge,value in [('top',55),('bottom',55),('left',75),('right',75)]:
        x=OxmlElement('w:'+edge); x.set(qn('w:w'),str(value)); x.set(qn('w:type'),'dxa'); margins.append(x)
    props.append(margins)
    for n,values in enumerate([headers]+list(rows)):
        r=t.rows[0] if n==0 else t.add_row()
        if n==0:
            repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); r._tr.get_or_add_trPr().append(repeat)
        else:
            # Ordinary records are short enough to keep together; main template
            # rows remain splittable so entire report chapters can span pages.
            no_split=OxmlElement('w:cantSplit'); r._tr.get_or_add_trPr().append(no_split)
        for i,(c,value) in enumerate(zip(r.cells,values)):
            c.width=Cm(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=c.paragraphs[0]; p.paragraph_format.line_spacing=1.15; p.paragraph_format.space_after=Pt(2)
            if i==0 and len(str(value))<18: p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            style_run(p.add_run(str(value)),font_size,n==0)
            if n==0:
                shade=OxmlElement('w:shd'); shade.set(qn('w:fill'),'E7E6E6'); c._tc.get_or_add_tcPr().append(shade)
    para(cell,'',size=3,indent=False,after=2)
    return t

def picture(cell,path,caption,width=15.3):
    path=Path(path)
    if not path.exists(): para(cell,'图件待生成 '+caption); return
    # Word can split even keepLines paragraphs inside the template's long outer
    # row. A single unsplittable nested row reliably binds heading, drawing and
    # caption while still allowing the surrounding chapter to flow across pages.
    headings=[]
    for child in reversed(list(cell._tc)):
        if child.tag==qn('w:p') and child.xpath('./w:pPr/w:outlineLvl'):
            headings.insert(0,child)
        else: break
    wrapper=cell.add_table(rows=1,cols=1); wrapper.autofit=False
    wrapper.columns[0].width=Cm(16); wrapper.alignment=WD_TABLE_ALIGNMENT.CENTER
    row=wrapper.rows[0]; row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    inner=row.cells[0]; inner.width=Cm(16); clear_cell(inner)
    for element in headings: cell._tc.remove(element); inner._tc.append(element)
    if headings:
        from docx.text.paragraph import Paragraph
        introduction={
            '图 4':'箭头表示模块调用依赖。四个入口共享核心库，文件系统适配与仓库承担统一的数据读写边界。',
            '图 5':'图中列出真实公开类型及其关联。命名空间函数不虚构为类；关键字段和行为见图后接口表。',
            '图 6':'顺序按时间向下展开。文件扫描先于编码与存储，只有完整清单提交后才向用户报告备份成功。',
            '图 7':'恢复先完成认证和数据校验，再写入临时文件；链接与目录元数据在普通文件数据之后处理。',
            '图 8':'客户端与服务器先建立认证加密会话，再协商缺失对象；对象完成后才提交可见快照索引。',
            '图 9':'计划时间和文件事件都进入同一任务状态机。失败进入重试等待，成功后才执行版本淘汰。',
        }
        text=next((v for k,v in introduction.items() if caption.startswith(k)), '下图对应本节设计，图后说明关键接口与执行约束。')
        h=Paragraph(headings[-1],inner); h.add_run().add_break(); style_run(h.add_run(text),11,False)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        b=OxmlElement('w:'+edge); b.set(qn('w:val'),'nil'); borders.append(b)
    wrapper._tbl.tblPr.append(borders)
    p=para(inner,'',indent=False,after=7); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.line_spacing=1.05
    p.add_run().add_picture(str(path),width=Cm(width)); p.add_run().add_break()
    style_run(p.add_run(caption),10)
    # Alternative text is attached to the actual drawing, not hidden body text.
    for prop in p._p.xpath('.//wp:docPr'): prop.set('descr',caption)

def bound_table_heading(cell,title,level=3):
    # The inherited long table ignores heading keepNext/pageBreakBefore. A
    # short nested block reliably binds the title and the following short table.
    wrapper=cell.add_table(rows=1,cols=1); wrapper.autofit=False
    wrapper.columns[0].width=Cm(16); wrapper.alignment=WD_TABLE_ALIGNMENT.CENTER
    wrapper.rows[0]._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    inner=wrapper.cell(0,0); inner.width=Cm(16); clear_cell(inner)
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        border=OxmlElement('w:'+edge); border.set(qn('w:val'),'nil'); borders.append(border)
    wrapper._tbl.tblPr.append(borders)
    heading(inner,title,level)
    return inner

def write_requirements(cases):
    lines=['# 澄匣备份需求与评分追踪','', '需求以课程 0.pdf 及两份教师补充说明为基准；分数由教师认定，本文不预记得分。','',
           '| 编号 | 能力 | 验收条件 | 主要模块 |','|---|---|---|---|']
    lines.extend('| '+' | '.join(r[:4])+' |' for r in REQS)
    lines += ['', '## 自主实现边界','', 'Qt/Win32 提供控件、文件与网络I/O、线程、系统事件和随机数。归档、压缩、密码原语及认证封装、快照增量、调度协议、纠错有限域算法由项目源码实现。DPAPI只保护当前用户本机凭据，不替代备份内容加密。', '',
    '## 必须保留的真实限制','', 'NTFS不映射Linux FIFO或设备节点。符号链接依赖系统权限；SKIP不得写作PASS。严格元数据权限不足失败，宽松模式告警。变化筛选基于元数据。通配解析和迭代匹配自主实现，支持* ? []范围及取反，使用逐码点Unicode simple fold；畸形集合和操作预算超限明确拒绝。纠错每组最多两个已定位损坏片，不代表任意20%损坏可恢复。网络预共享密钥协议没有前向保密。', '',
    '## 人工课程材料','', '姓名学号、成员提交、lg.team活动、证书、考勤、答辩参与和视频上传链接由小组真实补充。']
    (ROOT/'docs/requirements.md').write_text('\n'.join(lines)+'\n',encoding='utf-8')
    with (ROOT/'docs/traceability.csv').open('w',newline='',encoding='utf-8-sig') as f:
        writer=csv.writer(f); writer.writerow(['case_id','requirement_id','source_module','test_name','category','input_and_steps','expected','actual','duration_ms','evidence_file','report_section','demo_action'])
        for i,c in enumerate(cases,1):
            title,steps,expected=describe(c['name'])
            requirement=requirement_for(c['name']); req=next((r for r in REQS if r[0]==requirement),REQS[-2])
            writer.writerow([f'TC{i:03}',requirement,source_for(c['name'],req[3]),c['name'],c['category'],title+'；'+steps,expected,c['result']+('；'+c.get('detail','') if c.get('detail') else ''),c.get('durationMs',''),'evidence/test-results.json','软件测试报告 2 功能测试',req[1]])
        for filename,prefix,requirement,source,steps,expected in [
            ('gui-test-results.json','GUI','E16','src/gui/main_window.cpp','启动真实窗口，执行对应控件或备份/恢复/修复交互，检查状态与结果。','窗口渲染和功能断言均满足该测试定义。'),
            ('background-test-results.json','PROC','E17','src/apps/agent_main.cpp','启动独立Agent与GUI，按对应场景退出GUI、修改源或启动重复进程，检查任务和版本。','Agent持续运行、互斥或淘汰策略正确，生成版本可严格恢复。')]:
            for i,c in enumerate(load_json(ROOT/'evidence'/filename).get('cases',[]),1):
                writer.writerow([f'{prefix}{i:02}',requirement,source,c['name'],c['category'],steps,expected,c['result']+('；'+c.get('detail','') if c.get('detail') else ''),c.get('durationMs',''),'evidence/'+filename,'软件测试报告 7 GUI和部署复验', '真实窗口操作及GUI关闭后后台继续备份'])

def build():
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest()==EXPECTED_REFERENCE,'Original template changed'
    evidence,cases=choose_evidence(); by_name={c['name']:c for c in cases}; counts=Counter(c['result'] for c in cases)
    now=datetime.now(CHINA).strftime('%Y年%m月%d日 %H:%M')
    tested=evidence.get('date','未知'); qt=evidence.get('qt','见构建证据')
    write_requirements(cases)
    OUTPUT.mkdir(exist_ok=True); work=ROOT/'.cache/report-build'; work.mkdir(parents=True,exist_ok=True)
    doc=Document(str(REFERENCE)); outer=doc.tables[0]
    # Keep the original cover, source headers/footers/section and teacher region.
    for row_idx,row in enumerate(outer.rows):
        pr=row._tr.get_or_add_trPr()
        for item in list(pr):
            if item.tag==qn('w:cantSplit') or (4<=row_idx<17 and item.tag==qn('w:trHeight')): pr.remove(item)
    for idx in range(17):
        seen=set()
        for cell in outer.rows[idx].cells:
            if id(cell._tc) not in seen: clear_cell(cell); seen.add(id(cell._tc))
    def cell(i): return outer.rows[i].cells[0]
    para(outer.rows[0].cells[0],'小组 待小组填写',indent=False)
    para(outer.rows[0].cells[1],'组长与两名组员 姓名学号待小组填写',indent=False)
    for i,label in [(1,'组长'),(2,'组员一'),(3,'组员二')]:
        para(outer.rows[i].cells[0],label+'姓名 待小组填写',indent=False)
        para(outer.rows[i].cells[2],'学号 待小组填写',indent=False)
    heading(cell(4),'一 实验室名称',2); para(cell(4),'实验室与实际提交截止日期 待小组填写。报告技术证据生成时间 '+now+'。')
    heading(cell(5),'二 实验项目名称',2); para(cell(5),'澄匣备份 ChengVault Windows 数据备份软件',bold=True)
    heading(cell(6),'三 实验目的',2)
    para(cell(6),'本实验把 C++ 程序设计、数据结构与算法及软件工程方法应用到可运行的备份软件。课程训练重点是需求转化、模块设计、自主算法实现、可追溯测试和协同交付，而不仅是一个能够复制文件的界面。')
    para(cell(6),'项目围绕数据在备份与恢复过程中的可验证性设计：任何一次成功备份都应拥有完整清单；任何一次成功恢复都应通过数据校验；权限不足、取消和损坏必须得到明确结果。通过实现和对照测试理解文件系统语义、压缩编码、认证加密、增量存储、网络状态机和有限域纠错。')
    heading(cell(7),'四 实验内容',2)
    para(cell(7),'采用 C++20 与 Qt 6 Widgets 完成普通文件目录树备份恢复，覆盖课程全部扩展类别。源码独立编写；Aegis仅作为需求调研对象，不复制其代码、界面、报告或图表。系统使用两个归档格式、两种压缩和两种加密算法，另增加 RS 冗余纠错。')
    table(cell(7),['类别','本项目实现','验收说明'],[
    ['基础 40 分','目录树备份和指定位置恢复','比较文件内容、目录与空对象；不把文件复制误称为全部备份语义'],
    ['类型及元数据','NTFS 链接 稀疏 ADS 属主 时间 DACL','仅针对 NTFS；权限不足不得算成功'],
    ['自定义备份','六类必需筛选与四类附加筛选','路径 类型 名字 时间 尺寸 用户逐项验收'],
    ['打包 压缩 加密','索引/PAX Huffman/LZ77 AES256/ChaCha20','双向操作自主实现；加密覆盖文件名和清单'],
    ['自动与网络','定时 实时 淘汰 网盘 用户 元数据 传输 增量','真实后台进程和 TCP，会话隔离及断点续传'],
    ['新增扩展','RS12 10 冗余修复','每组两片已定位损坏能力，超过能力明确失败']], [2.7,7.0,6.3])
    heading(cell(8),'五 实验器材和工具',2)
    table(cell(8),['项目','实际环境'],[
    ['处理器与内存','Intel Core Ultra 9 285H，16 核 16 逻辑处理器，物理内存约 31.59 GiB'],
    ['操作系统','Windows 11 Version 25H2 x64；测试数据位于 NTFS 卷'],
    ['语言与框架',f'C++20，Qt {qt}，Qt Core Widgets Network Concurrent Test'],
    ['编译构建','隔离 MSYS2 UCRT64 GCC 16.2.0，CMake 与 Ninja；实际版本记录以构建日志为准'],
    ['设计和协作','可编辑 StarUML 模型、Git 分模块提交；lg.team 由小组维护真实任务和成员记录'],
    ['文档与验证','Microsoft Word 原生 PDF 导出；Python 生成模板副本；Poppler 渲染逐页核验']], [3.0,13.0])
    heading(cell(9),'六 实验步骤及操作',2).paragraph_format.keep_with_next=False
    c=cell(10); heading(c,'需求分析说明书（10分）',1)
    heading(c,'1 任务概述',2); heading(c,'1.1 引言',3)
    para(c,'澄匣备份用于将个人或实验项目目录保存为可校验的独立归档或多版本仓库，并在原数据丢失、误改或需迁移时恢复。本报告面向任课教师和项目维护者，依次给出需求、实现依据、设计模型、可重现测试及当前限制，便于把评分项对应到真实代码和证据。')
    para(c,'需求来源为课程 0.pdf、老师发的东西.txt、其他评分细则或者注意事项.txt 和指定实验报告模板。基础要求是目录树数据备份与还原；本组选定全部扩展类别。使用第三方库直接完成计分功能会影响扩展分，因此算法与业务模块由项目实现，Qt 和 Win32 只承担基础设施。')
    heading(c,'1.2 综合描述',3)
    para(c,'主要参与者是备份使用者和服务器管理员。备份使用者选择源、策略、密码与目的地，查看版本、执行恢复和检查损坏；服务器管理员维护账户、设备授权和配额。系统时钟与 Windows 文件事件是自动任务的触发来源，不是具有业务目标的人类角色。')
    para(c,'两种使用方式共用文件系统核心。独立归档便于拷走一个文件，配套可选纠错文件；版本仓库把文件拆为 1 MiB 块，用完整快照清单组织共享内容，便于重复备份和网络增量。选择密码加密后，文件内容与清单均加密。登录密码、仓库密码、设备配对密钥相互独立。')
    heading(c,'1.3 运行环境与边界',3)
    para(c,'目标为 Windows 10/11 x64 和 NTFS。程序可作为普通用户运行，但符号链接创建、非当前用户属主及限制性 ACL 的恢复取决于系统授权。课程列举的管道和设备是平台示例，本项目不宣称将 Linux FIFO 或设备节点还原成 NTFS 同类对象。未知重解析点明确拒绝，避免误备份云占位对象。')
    para(c,'文件正在被其他进程写入时，捕获采用拒绝共享写入的读句柄；冲突会失败并由任务重试。项目没有 VSS，不声称数据库跨文件事务一致性；目录整体恢复也不是全局回滚事务。文件级覆盖使用验证后的同目录临时对象替换，旧文件不会预先删除。')
    heading(c,'2 功能需求',2); heading(c,'2.1 功能划分和验收编号',3)
    para(c,'下表将课程计分要求转为可复验的项目需求编号。编号在源代码模块、测试 CSV 和答辩演示中保持一致，具体得分由教师依据实现质量判定。')
    table(c,['编号','需求','可观察验收条件'],[(r[0],r[1],r[2]) for r in REQS],[1.4,3.4,11.2])
    heading(c,'2.2 用户交互方式',3)
    para(c,'主窗口使用侧边导航组织总览、备份、恢复、自动任务、远程仓库、归档修复、日志和设置。备份流程依次选择源、预览筛选、配置归档或仓库、选择算法和密码、启动任务；恢复先打开归档或仓库并完成认证，再选择版本、路径和目的地。耗时工作进入后台线程，界面通过 Context 接收进度与取消。')
    para(c,'界面展示文件与任务信息。筛选对话框通过路径、类型、名字、时间、尺寸、用户及额外条件的字段配置规则，并可导入旧清单用于变化筛选；预览显示匹配与原因。命令行以 JSON 计划表达相同规则。错误提示区分配置错误、损坏、权限和取消；元数据警告保留在结果中供核查。')
    picture(c,ROOT/'evidence/gui/backup.png','图 1 备份界面及策略设置')
    picture(c,ROOT/'evidence/gui/restore.png','图 2 恢复界面及目标保护选项')
    heading(c,'2.3 系统用例图与规范说明',3)
    para(c,'用例图的系统边界为 ChengVault，参与者在边界外，用例名称采用用户能够理解的动作。备份和恢复是核心用例，查看版本、验证归档、配置任务、远程同步和修复归档围绕其形成独立目标。图中关系与以下用例编号一致；密码认证、数据校验等共同步骤在流程描述中说明，不把每个内部函数画成独立业务用例。')
    picture(c,ROOT/'docs/uml/01-use-cases.png','图 3 系统用例图')
    use_cases=[
    ('UC01 创建备份','备份使用者','源可读，目标可写；加密时有口令','选择源→预览筛选→选择格式/压缩/加密→开始→查看结果','源变化或被占用则失败；取消不发布快照；成功后可检验归档或快照。'),
    ('UC02 恢复数据','备份使用者','归档或快照存在，必要口令正确','打开并认证→选择版本和文件→选择目标→预检→恢复','默认拒绝覆盖；块损坏先失败；严格元数据失败需处理权限。'),
    ('UC03 管理自动任务','备份使用者','仓库已初始化，当前用户可保存 DPAPI 凭据','选择间隔/每日/每周或实时→设置淘汰→启用后台','错过计划会补跑；同时任务互斥；失败重试并记录错误。'),
    ('UC04 远程同步','备份使用者','服务端可达，账户和设备已授权','登录认证→上传缺失对象→提交索引→下载并验证','断线保留暂存偏移；账户隔离；非空下载目标拒绝。'),
    ('UC05 检查与修复归档','备份使用者','原归档或受损归档及对应 .cvr 存在','生成冗余→检查损坏→写新文件修复→完整摘要验证','每组超过两片损坏或索引冲突拒绝；旧文件保留；可单独验证归档。'),
    ('UC06 管理账户','服务器管理员','管理员通过设备及密码认证','创建或禁用账户→重置口令/设配额→注册或撤销设备','普通账户被服务端拒绝；列表不泄漏密码散列和密钥。'),
    ('UC07 按规则扫描','备份使用者，由创建备份包含','至少有一个源','设置六类/额外规则→预览命中与原因→保存策略','非法时间拒绝；排除优先；命中文件祖先目录保留。'),
    ('UC08 校验并安全恢复','备份使用者，由恢复数据包含','加密对象已认证，清单可读取','验证安全路径和全部数据块→临时写入→校验→提交→恢复元数据','错误密码/截断/篡改预先失败；已有目标默认保护；严格元数据失败明确报告。'),
    ('UC09 用户设备认证','备份使用者，由远程同步包含','已注册设备配对密钥和有效账户','双方随机数→认证握手→派生方向密钥→加密发送登录口令','错误设备密钥、禁用账户或错误密码被拒绝；重放帧不执行。'),
    ('UC10 定位损坏分片','备份使用者，由检查修复包含','至少一份修复元数据副本有效','验证索引副本→按片计算SHA256→标记损坏位置→判断可修范围','元数据两份均坏或冲突拒绝；数据和校验片共同计数，同组超过两片拒绝。')]
    for uc,actor,pre,flow,exception in use_cases:
        heading(c,uc,3)
        table(c,['要素','描述'],[['主要参与者',actor],['前置条件',pre],['基本流程',flow],['异常与后置条件',exception]], [2.4,13.6])
    heading(c,'3 外部接口需求',2)
    table(c,['接口','输入输出与约束'],[
    ['GUI','Qt Widgets 中文表单；任务结果包括成功、文件数、逻辑字节、存储字节、警告与结构化详情。'],
    ['CLI','独立客户端命令支持 backup restore verify repo-backup repo-restore 及网络/修复操作；失败返回非零退出码；密码可通过文件提供。'],
    ['文件系统','Win32 UTF16 扩展长度路径，清单 UTF8；禁止归档路径穿越与保留设备名。'],
    ['网络','TCP 默认 127.0.0.1:48484；局域网需显式监听地址及防火墙配置；认证帧长度有上限。'],
    ['本机凭据','后台主密钥由当前用户 DPAPI 保护；不能直接迁移到另一 Windows 账户解封。']], [2.8,13.2])
    heading(c,'4 非功能需求',2)
    para(c,'正确性优先于静默继续。归档长度、资源数量、流名称、块偏移和数据摘要都必须验证；不存在以空数据替代读取失败的路径。错误密码不能创建恢复目录，取消不能发布未完成快照。日志不得写入密码、主密钥或设备密钥。')
    para(c,'可用性要求同一核心接口被 GUI、CLI 和后台复用。核心使用有界块处理大文件，不一次载入整份归档；独立归档转换需要临时磁盘空间，PAX 对稀疏洞会先展开逻辑零字节。性能数据须记录真实测试环境、数据集、耗时和峰值内存，不给未经测量的速度保证。')
    para(c,'可维护性要求文件系统、归档、压缩、密码、仓库、网络、调度和纠错分开，命名统一，注释解释算法和边界。使用 CMake 与脚本自动构建，Git 按模块留下多次真实提交。白盒测试覆盖算法标准向量与异常分支；源代码覆盖率目标为行 85% 和分支 75%，实测未达到时如实列出差距。')
    heading(c,'5 项目规划和分工',2)
    para(c,'下表分配现有模块的接手、维护与答辩责任，不代表历史作者。身份与实际贡献待本人据实填写；精确文件及提交流程见《ChengVault三人分工与提交指南》和 docs/team-work-plan.csv。')
    table(c,['成员身份','现有模块责任及接手内容','交叉复核'],[
    ['A 姓名学号待填','GUI/CLI、NTFS、筛选、两套归档及配套测试；统筹集成、需求设计和界面图文','B 复核 A'],
    ['B 姓名学号待填','压缩、密码、RS 纠错、公共安全与算法测试；复验标准向量、78 种损坏组合和原始覆盖率','C 复核 B'],
    ['C 姓名学号待填','版本仓库、后台调度、网络及配套测试；统筹构建发布、演示和两机验收','A 复核 C']], [3.0,9.7,3.3])
    para(c,'三人各自读码、测试、真实修改并提交，均须参加答辩。',size=10,after=2)

    c=cell(11); heading(c,'系统设计文档（20分）',1)
    heading(c,'1 开发环境和工具',2); heading(c,'1.1 开发环境',3)
    para(c,f'工程使用 C++20 和 Qt {qt}，分为 ChengVault 图形客户端、ChengVaultCli 命令行、ChengVaultAgent 后台以及 ChengVaultServer 服务端。共享静态核心库避免各入口重复实现业务。CMake 自动收集模块，Ninja 编译；依赖目录和发布目录与源码分离。')
    heading(c,'1.2 依赖库和自主实现边界',3)
    table(c,['能力','允许的基础设施','本项目自主逻辑'],[
    ['界面与文件','Qt Widgets Core，Win32 文件 API','策略、路径检查、元数据清单、恢复提交'],
    ['算法','系统安全随机数；普通字节容器','Huffman LZ77 SHA256 HMAC PBKDF2 HKDF AES256 ChaCha20'],
    ['网络与调度','Qt TCP、线程、定时器，Windows 文件通知','帧协议 认证 序号 用户隔离 增量 续传 调度 淘汰'],
    ['纠错与凭据','普通文件 I/O，DPAPI 本机凭据封装','GF256 运算 RS 系统化编码 矩阵求逆 分片修复']], [2.7,5.8,7.5])
    para(c,'运行时没有使用压缩库、归档库、OpenSSL 或现成网盘组件完成计分功能。独立 TAR 程序仅用于测试格式兼容性，不参与用户备份流程。DPAPI用于保护本机持久化凭据；备份内容仍由自主 AES256/ChaCha20 认证封装保护。')
    heading(c,'2 总体设计',2); heading(c,'2.1 系统结构设计',3)
    picture(c,ROOT/'docs/uml/02-components.png','图 4 系统构件与接口')
    para(c,'上层入口只负责参数、交互和任务启动，核心库定义 BackupPlan、RestoreOptions、TaskResult 和 Context。文件系统模块对外输出与读取统一清单；仓库模块决定块的稳定标识与加密存储；归档模块决定独立文件的布局。该边界使网络不需要理解明文文件，纠错也可以直接作用于最终密文归档。')
    table(c,['构件','职责与接口','失败约束'],[
    ['文件系统','preview capture restore validateManifest；ChunkSink/ChunkSource 交换原始块','读取失败不伪造内容，路径与流校验先行'],
    ['归档','indexed 与 PAX 内层；通用压缩加密外层','完整终结记录验证后再恢复'],
    ['密码和压缩','独立算法、标准向量、大小上限','先认证后明文，拒绝计数器耗尽'],
    ['仓库','初始化、备份、列举、恢复、校验、换密码、淘汰','锁保护清单，成功清单最后提交'],
    ['调度','nextRun、DirectoryWatch、Engine；DPAPI 任务密钥','互斥运行、重试和事件合并'],
    ['网络','Client 与 Server；管理、传输和会话','用户隔离、配额、序号、原子索引'],
    ['纠错','创建侧车、检查、恢复到新文件','摘要确认后才发布，超过能力失败']], [2.5,7.6,5.9])
    heading(c,'2.2 关键设计取舍',3)
    para(c,'选择固定 1 MiB 块，是为获得简单可审计的偏移和增量边界。局部修改通常只新增对应块；文件头插入可能使后续固定块发生变化，这是固定分块的已知代价。每个快照保存完整清单，因此恢复无需依次应用历史增量。去重在仓库内进行，跨用户或跨仓库不共享对象。')
    para(c,'单文件索引归档去除重复块并建立定位表；PAX 归档保存真实文件、目录和链接成员，具有独立读取器可检查的格式结构。两者共用外层封装只是复用压缩与加密，并不把两种内层布局混为一种。PAX 内部稀疏文件洞展开为零，索引格式更适合极大稀疏数据。')
    heading(c,'3 静态建模',2)
    picture(c,ROOT/'docs/uml/03-core-classes.png','图 5 核心类型和类关系')
    para(c,'BackupPlan 聚合源列表、目标、格式、压缩算法、密码算法与 filters；RestoreOptions 保存覆盖、严格元数据和选择路径。Context 的原子取消标志和进度回调在边界检查点被调用，不与 GUI 控件耦合。Repository 拥有稳定主密钥、配置及存储路径；Client 拥有连接参数，Server 管理有界会话。Engine 组合文件监听器并组织后台任务。')
    table(c,['接口或类型','数据与行为','设计约束'],[
    ['BackupPlan','sources destination format compression cipher password filters strictMetadata','默认序列化省略密码'],
    ['RestoreOptions','overwrite strictMetadata selectedPaths','默认 overwrite=false'],
    ['TaskResult','success files bytes storedBytes warnings details','逻辑数据量与存储量分开'],
    ['Context','cancelled 与 progress(done,total,text)','取消在安全检查点抛结构化错误'],
    ['ChunkSink ChunkSource','原始 QByteArray ↔ 块 ID','单块最多 1 MiB；ID 由存储层决定'],
    ['Repository','backup restore list manifest prune verify changePassword','明文密钥只保留在当前进程'],
    ['RemoteOptions','主机 端口 用户 登录密码 设备 ID 配对密钥','不把备份密码发送给服务器']], [3.5,7.2,5.3])
    heading(c,'4 动态建模',2); heading(c,'4.1 备份顺序',3)
    picture(c,ROOT/'docs/uml/04-backup-sequence.png','图 6 备份顺序图')
    para(c,'备份先验证源目标关系并扫描元数据，再执行过滤和祖先闭包。每个被选中流以稳定读句柄读取，分配区间中的块交给存储层；普通文件的大小和最后写入时间必须仍与扫描一致。归档模式先生成内层，再压缩与认证加密；仓库模式先持久化新增块，最后提交完整快照。任何异常均不能返回一个未完成的成功快照。')
    para(c,'通配符筛选由项目自编解析器和最近星号迭代回退实现，支持 *、?、字符集合、区间和取反。? 消耗一个原 Unicode 码点，逐码点 simple fold 保持大小写不敏感且不把 ß 展开成 SS。无通配路径规则按分隔符边界包含子树。畸形集合在扫描前拒绝，单条规则最多 32000 个 UTF16 单元，匹配操作设 8 Mi 次上限；最坏 O(P×T)、空间 O(P+T)。Qt 只提供字符容器和大小写转换，不调用正则或通配匹配引擎。')
    heading(c,'4.2 恢复顺序',3); picture(c,ROOT/'docs/uml/05-restore-sequence.png','图 7 恢复顺序图')
    para(c,'恢复先验证密码、封装和清单，检查所有需要的数据块。随后创建目录、写普通文件和 ADS、建立链接，最后恢复目录时间与 ACL。选择硬链接时自动包含其数据目标；选择子文件时保留祖先。已有目标默认拒绝，显式覆盖时只在临时文件完成且校验通过后同卷替换。')
    heading(c,'4.3 网络顺序',3); picture(c,ROOT/'docs/uml/06-network-sequence.png','图 8 网络认证与增量传输')
    para(c,'客户端与服务端交换随机数，使用绑定账户的设备 PSK 认证握手，再以 HKDF 派生不同方向密钥。登录口令在加密通道内验证。上传先提交 inventory 并得到缺失对象及暂存偏移，逐段上传后核验密文摘要，commit 才替换远端索引；旧索引在失败期间继续有效。下载验证全部对象后发布仓库，再由客户端口令解密恢复。')
    heading(c,'4.4 自动任务状态',3); picture(c,ROOT/'docs/uml/07-task-state.png','图 9 自动备份任务状态')
    para(c,'任务从等待进入就绪，再进入运行；成功后更新下次时间并执行淘汰，失败记录原因后等待重试，取消由检查点退出。文件通知线程只报告脏状态，不直接备份；Engine 对事件合并并等待稳定窗口。目录监听退出时取消挂起 I/O，避免后台进程因等待事件而无法关闭。')
    heading(c,'5 数据设计',2); heading(c,'5.1 清单和文件系统语义',3)
    table(c,['结构','主要字段','校验规则'],[
    ['清单根','version createdUtc entries fileCount totalBytes warnings','版本为 1，entries 有数量上限'],
    ['文件条目','path type size metadata streams target','相对路径、大小写无冲突、父项必须是目录'],
    ['数据流','name size sparse ranges chunks','默认流名为空；ADS 形如 :name:$DATA'],
    ['块引用','id offset length sha256','偏移递增不重叠，长度 1 至 1 MiB'],
    ['元数据','FILETIME 字符串 attributes ownerSid groupSid daclBase64 securityControl','时间避免浮点失真；SID 与 ACL 结构有效'],
    ['链接','target targetDirectory reparseBase64','硬链接直指普通文件；重解析标签限定支持范围']], [2.7,7.2,6.1])
    para(c,'硬链接按卷号和文件 ID 识别，只保存首个选中路径的数据；恢复后各路径必须指向同一文件 ID。稀疏流保存逻辑长度及分配区间，只写有数据的块。符号链接保留目标文本；绝对链接仍指向原绝对目标，系统不擅自重定向。目录和文件命名数据流都纳入清单和块引用，不能仅处理主流。')
    para(c,'属主和组使用 SID 持久化，显示账户名仅供交互。权限保存 DACL 原始 ACE 及保护/继承状态。恢复利用 MAXIMUM_ALLOWED 句柄的 SetSecurityInfo 行为抑制父目录 ACE 向子对象再次传播，确保每个子对象的独立记录不会被最后的父目录恢复覆盖。SACL 审计、EFS 私钥和 NTFS 压缩状态不在承诺范围。')
    heading(c,'5.2 两种归档和共同封装',3)
    table(c,['布局','字段和编码','作用'],[
    ['CVARCH2 外层','魔数 头长度 JSON；每记录含序号 类型 头摘要 数据；终结记录含总长与全流摘要','逐 1 MiB 压缩/加密，检测重排 截断 追加'],
    ['索引内层','CVIDX01、清单长度/JSON、唯一块数、SHA256 ID/长度/块字节','只存唯一块；读取时建立偏移表'],
    ['PAX 内层','512 字节 USTAR 头、PAX path/size/linkpath、真实文件数据、零填充','支持长路径与大长度；独立 tar 可读'],
    ['PAX NTFS 扩展','首成员清单；CHENGVAULT.kind/path/stream 键','附加流单独成员，完整 NTFS 语义由清单恢复']], [3.1,8.7,4.2])
    para(c,'外层多字节整数显式小端编码，网络帧长度使用大端，不能直接把 C++ 结构体写入文件。解压输出受 1 MiB 上限约束，外层信封不超过 4 MiB，清单受 64 MiB 上限约束。加密归档的主密钥随机生成，口令通过带盐 120000 次 PBKDF2 派生封装密钥；每个对象再次随机派生加密与认证密钥。')
    heading(c,'5.3 压缩和密码算法',3)
    para(c,'Huffman 先统计符号频数，建立最小权重合并树并为叶子生成前缀码；解码根据相同频数模型恢复，覆盖空输入、单符号和多符号位串。LZ77 在有限滑动窗口中寻找最长匹配，输出字面量与距离/长度；重叠复制逐字节推进，零距离与越界回溯拒绝。压缩后若不节省空间，使用带原始块标记的数据。')
    para(c,'AES256 执行 256 位密钥扩展与轮变换，以 CTR 模式产生密码流；ChaCha20 执行加法、异或、循环移位组成的四分之一轮和 20 轮状态变换。二者的密码流异或是标准算法步骤，不是以单一 XOR 冒充加密算法。信封使用 Encrypt then MAC；HMAC 认证通过前不得返回明文，计数器用尽拒绝继续。')
    para(c,'SHA256、HMAC、PBKDF2 与 HKDF 同样自主实现。验证采用 FIPS、NIST 和 RFC 的固定已知答案，避免只用同一错误实现的加密解密相互抵消来判断正确。随机数由 Windows 提供，主密钥不由文件名或时间戳派生。课程实现没有经过独立密码审计，报告不把它描述为商用密码认证产品。')
    heading(c,'5.4 仓库与网络元数据',3)
    table(c,['对象','内容','一致性与隐私'],[
    ['repository.json','仓库 ID、算法、随机盐、口令封装主密钥','换密码只改封装，不重写密文数据'],
    ['快照','完整文件清单及版本摘要','客户端主密钥加密；服务器不可读文件名'],
    ['数据块','仓库专属 HMAC 标识与压缩认证密文','同仓库复用，跨用户不共享'],
    ['远端 index','账户名下仓库的密文对象路径 大小 摘要','commit 原子替换，失败保留旧索引'],
    ['账户及设备','口令盐/派生值 角色 配额 设备 DPAPI 密钥','服务端认证；不储存备份主密钥']], [3.0,7.5,5.5])
    para(c,'网络协议的最大帧为 8 MiB，每个传输段最多 1 MiB，对象最大 64 MiB，inventory 受数量与帧长双重上限。严格检查会话标识和递增序号，拒绝重放与跨会话数据。预共享密钥设计便于本机及局域网演示，但没有前向保密；配对密钥必须通过可信方式交付，不能从未认证 TCP 自动取得。')
    para(c,'配额在提交互斥区再次检查，避免并发提交同时耗尽余额。服务端按账户路径隔离存储，旧索引不再引用的 blob 在成功提交后清理。中断上传暂存支持续传，当前没有自动过期清理策略；管理员应监测暂存占用。客户端锁住本地仓库直至上传完成，避免传输过程中清单被另一任务修改。')
    heading(c,'5.5 RS 冗余纠错设计',3)
    para(c,'新增纠错功能作用于已经完成压缩和加密的归档字节。每组 10 个 1 MiB 数据分片加 2 个校验分片，在 GF256 中使用系统化 Vandermonde 矩阵编码；每片 SHA256 定位损坏，剩余任意十片组成可逆子矩阵，通过高斯消元重建缺失数据。有限域乘法、矩阵求逆及编码循环均由项目编写。')
    para(c,'侧车保存原归档长度、编码参数、分片摘要和整包摘要，头尾各有一份校验元数据。至少一份有效且两份不冲突才能修复。末组零补齐后按原长度截断；输出写入新文件并通过整包摘要后发布。每组最多两个已定位损坏片包括数据片和校验片，不能解释为任意位置损坏 20% 都可恢复。完整分组校验数据开销 20%，小文件末组及索引使实际比例增加。')
    picture(c,ROOT/'evidence/gui/repair.png','图 10 归档检查与冗余修复界面')

    c=cell(12); heading(c,'软件测试报告（20分）',1)
    heading(c,'1 引言和测试基线',2)
    para(c,f'测试基线读取 evidence/test-results.json，执行时间标识为 {tested}。本次记录 {len(cases)} 个测试：通过 {counts["PASS"]}，失败 {counts["FAIL"]}，跳过 {counts["SKIP"]}。下列表格由实际 JSON 生成，不把测试代码的存在当作已通过。完整记录和需求映射见 docs/traceability.csv；每个用例 ID 对应一个真实测试名。')
    para(c,'测试采用项目自建 C++ 断言运行器，异常使该用例 FAIL，SKIP 单独计数。算法测试使用公开标准已知答案；集成测试运行真实 Win32 对象、真实 TCP 套接字及临时数据，不通过模拟返回值假装网络连接。部分写入中途故障用回调注入，明确标注而不宣称真实填满磁盘。')
    categories=Counter(x['category'] for x in cases)
    category_focus={'白盒单元':'算法分支、标准向量、边界与状态转换','安全':'认证、权限、路径、长度、用户隔离','集成':'核心模块组合与真实系统 API','系统功能':'完整备份恢复和特定类型','性能':'复用数据与大文件边界','异常健壮性':'截断、损坏、取消和故障清理','并发稳定性':'任务锁、取消和监听退出','部署兼容性':'独立格式读取器与运行环境','功能':'计划时刻与业务结果','边界健壮性':'非法/极值参数'}
    table(c,['测试类别','记录数','验证重点'],[(k,v,category_focus.get(k,'按实际类别对应源码断言')) for k,v in categories.items()],[3.0,1.8,11.2])
    heading(c,'2 功能测试与规范用例记录',2)
    para(c,'公共前置条件：运行 Release 构建的 ChengVaultTests，测试盘为 NTFS，测试程序可建立自身临时目录；网络用例使用随机可用本机端口和专用测试账户。正文选取各模块代表性用例及本轮修复回归，每行给出输入、操作和具体期望；全部执行记录与描述保存在 CSV，不将组合展开和重复执行扩大为新用例。需要权限的符号链接用例在权限不足时跳过，不能计入通过率分子。')
    selected=[]
    for i,test in enumerate(cases,1):
        name=test['name']
        # The 78 RS matrix cases are summarized separately to avoid disguising
        # a repeated combinatorial matrix as broad product scenario coverage.
        if name.startswith(('RS12-10/single-erasure-','RS12-10/double-erasure-')): continue
        if name.startswith('SHA256/stream-split-') and not name.endswith(('0','63','64','65')): continue
        if name.startswith('Envelope/') and '/roundtrip-' in name and not name.endswith(('0','65')): continue
        if name.startswith('Envelope/') and '/tamper-' in name and not name.endswith(('-0','-1')): continue
        if name.startswith('Compression/') and '/fixture-' in name and int(name.rsplit('-',1)[1])>3: continue
        if name.startswith(('Huffman/roundtrip-', 'LZ77/roundtrip-')) and not name.endswith(('-0','-1','-65','-4096')): continue
        selected.append((i,test))
    assert len(selected)>=96, f'Only {len(selected)} detailed executed cases'
    groups=[('2.1 密码及基础算法',lambda n:n.startswith(('SHA','HMAC','PBK','HKDF','AES','ChaCha','KDF','Crypto','Envelope','Entropy','Credentials','HashFile','Common/'))),
            ('2.2 压缩和解压',lambda n:n.startswith(('Compression','Huffman','LZ77'))),
            ('2.3 纠错和损坏恢复',lambda n:n.startswith(('RS12','Repair'))),
            ('2.4 文件系统和筛选',lambda n:n.startswith('fs.')),
            ('2.5 归档格式组合',lambda n:n.startswith('archive.')),
            ('2.6 版本仓库',lambda n:n.startswith('repository.')),
            ('2.7 网络服务',lambda n:n.startswith('network.')),
            ('2.8 自动任务与事件',lambda n:n.startswith('SCH-'))]
    covered=set()
    priorities=('pax-padding-nonzero','junction-print-name','display-target-mismatch','chunk-changes-after-preflight',
        'fs.roundtrip','fs.restore.no-overwrite','fs.restore.explicit-overwrite','hardlinks-shared-identity',
        'file-and-directory-ads','junction-no-recursion','reparse-escape','archive.roundtrip.',
        'archive.wrong-password','archive.pax-independent-reader','archive.truncation-detected',
        'disk-full-write-boundary','denied-dacl','owner-group-times','symbolic-file-and-directory','sparse-over-4gib',
        'path-exclusion','object-type','owner-account','time-bounds','size-inclusive','content-boundary','changed-since',
        'glob.one-codepoint','glob.simple-fold-sharp-s','glob.digit-range','glob.negated-set','glob.invalid-[z-a]',
        'glob.literal-directory','glob.explicit-work-budget','pax-alias-ntfs','cipher-alias','cancel-during-wrap',
        'AES256CTR/NIST','ChaCha20/RFC','HKDF/RFC','independent-parameter-lengths','partial-block-and-HKDF','two-stripes-independent','whole-digest-inconsistent',
        'reconstructed-shard-hash','source-changed-length','concurrent-output-preserved','legitimate-overlapping',
        'intrasnapshot-dedup','authenticated-manifest','corrupt-retained','age-retention','list-lock-protects','download-cancel-resume',
        'destination-changed','same-hash-different','session-invalidated','session-revoked','connection-limit',
        'SCH-18','SCH-19','SCH-21','SCH-25','SCH-26','SCH-31','SCH-32','SCH-33','SCH-34','SCH-36','SCH-40','SCH-42','SCH-43','SCH-44')
    limits=(28,16,22,48,32,19,27,28)
    for (title,predicate),limit in zip(groups,limits):
        candidates=[x for x in selected if predicate(x[1]['name'])]
        important=[x for x in candidates if any(k in x[1]['name'] for k in priorities)]
        regular=[x for x in candidates if x not in important]
        chosen=sorted((important+regular)[:limit],key=lambda x:x[0])
        rows=[]
        for i,t in chosen:
            covered.add(i); feature,steps,expected=describe(t['name'])
            actual={'PASS':'通过','FAIL':'失败','SKIP':'跳过'}.get(t['result'],t['result'])
            if t.get('durationMs') is not None: actual+='\n'+str(t['durationMs'])+' ms'
            if t.get('detail'): actual+='\n'+t['detail'][:95]
            rows.append([f'TC{i:03}\n{requirement_for(t["name"])}',feature,steps+'\n预期 '+expected,actual])
        if rows:
            first_table_cell=bound_table_heading(c,title)
            # Nested table headers are not repeated by Word across a splitting
            # outer-table row. Short tables repeat the header explicitly.
            for start in range(0,len(rows),8):
                table(first_table_cell if start==0 else c,['编号','场景及输入','操作步骤和预期结果','实际'],rows[start:start+8],[1.5,3.6,8.7,2.2],9.5)
    assert len(covered)>=96
    heading(c,'2.9 RS 组合完备性',3)
    rs=[x for x in cases if x['name'].startswith(('RS12-10/single-erasure-','RS12-10/double-erasure-'))]
    para(c,f'RS 矩阵测试另有 {len(rs)} 个实际组合记录。目标为 12 个单片位置和 C(12,2)=66 个双片组合；每个组合从同一固定数据集编码，删除指定片后解码并逐片比较。组合矩阵验证任意合法位置，而前面的损坏归档、元数据副本与错误侧车用例验证文件层边界。')
    rsrows=[]
    for i in range(12):
        one=by_name.get(f'RS12-10/single-erasure-{i}',{})
        pairs=[by_name.get(f'RS12-10/double-erasure-{i}-{j}',{}) for j in range(i+1,12)]
        rsrows.append([i,one.get('result','未执行'),str(len(pairs)),', '.join(f'{i}+{j}' for j in range(i+1,12)) or '无后续组合','全部 PASS' if pairs and all(x.get('result')=='PASS' for x in pairs) else '不适用' if not pairs else '见失败记录'])
    table(c,['片号','单片','双片数','同组组合','实际'],rsrows,[1.1,1.9,1.5,8.7,2.8],9.5)
    heading(c,'3 代码测试和白盒依据',2)
    para(c,'白盒重点包括 Huffman 的空输入/单符号分支、LZ77 重叠复制与非法距离、AES/ChaCha 计数器极限、HMAC 标签失败、PAX 长度和校验、清单路径闭包、仓库锁与提交、调度时间边界、RS 奇异/缺片条件。测试根据外部标准答案或可观察不变量进行断言，避免逐行照抄实现逻辑。')
    quality=load_json(ROOT/'evidence/quality/summary.json')
    if quality:
        para(c,f'原创生产源码及头文件注释按物理行统计为 {quality.get("commentPercentPhysical")}%（不计测试源码）；格式检查 {quality.get("formatCheckedFiles")} 个文件，不合规 {len(quality.get("formatNoncompliantFiles",[]))} 个。cppcheck 记录 error {quality.get("cppcheckErrorCount")}、warning {quality.get("cppcheckWarningCount")}，另有 {quality.get("cppcheckFindingCount")} 条已审查的风格/性能建议，原始 XML 与评审说明保留，不能称为完全无提示。')
    cov=load_json(ROOT/'evidence/coverage/summary.json')
    if 'line_percent' in cov:
        para(c,f'覆盖率范围为全部 src/core 与 src/platform 源文件，实际证据 evidence/coverage/summary.json。行覆盖 {cov["line_covered"]}/{cov["line_total"]}={cov["line_percent"]}%；GCC 全控制流分支覆盖 {cov["branch_covered"]}/{cov["branch_total"]}={cov["branch_percent"]}%，其中包含编译器生成的异常处理路径。行目标为 85%，分支目标为 75%；当前结论分别为 '+('达到' if cov['line_percent']>=85 else '未达到')+'、'+('达到' if cov['branch_percent']>=75 else '未达到')+'。不同口径不能混合或用辅助百分比替代原目标。')
        table(c,['核心模块','行覆盖','分支覆盖'],[(Path(x['filename']).name,f'{x["line_percent"]}%',f'{x["branch_percent"]}%') for x in cov.get('files',[])],[7.0,4.5,4.5],9.5)
        auxiliary=load_json(ROOT/'evidence/coverage/auxiliary-summary.json')
        if 'branch_percent' in auxiliary:
            para(c,f'辅助统计 evidence/coverage/auxiliary-summary.json 排除编译器异常/不可达路径后分支覆盖为 {auxiliary["branch_percent"]}%。该数值解释异常密集 C++ 代码的统计差异，仍与上面的全控制流统计并列公开，不替换原始门槛结论。')
        if 'decision_percent' in cov:
            para(c,f'同份原始 gcovr 报告另列判定覆盖 {cov["decision_covered"]}/{cov["decision_total"]}={cov["decision_percent"]}%。判定、全控制流分支和排除异常后的辅助分支三个口径分别报告；不能把其中最高值冒充原定分支指标。')
    else:
        para(c,'行覆盖目标 85%、分支目标 75%。生成时尚无最终 gcov 汇总，不能填写假定百分比。')
    heading(c,'4 性能测试',2)
    perf_files=sorted((ROOT/'evidence/benchmark').glob('*/benchmark-results.json')) if (ROOT/'evidence/benchmark').exists() else []
    perf=load_json(ROOT/'evidence/benchmark-results.json') or (load_json(perf_files[-1]) if perf_files else {})
    if perf:
        para(c,f'性能脚本实际数据规模 {perf.get("dataMiB","未记录")} MiB，{perf.get("smallFileCount","未记录")} 个小文件；运行标识 {perf.get("timestampUtc","")}。采用同机器同脚本比较，峰值工作集为进程测量值；短进程采样可能为 0，不能解释为程序不占内存。')
        rows=[]
        for rec in perf.get('records',[]):
            result=rec.get('result')
            actual=('成功' if result.get('success') else '失败') if isinstance(result,dict) else str(result or ('PASS' if rec.get('exitCode')==0 else rec.get('error','见日志')))
            if isinstance(result,dict) and result.get('warnings'): actual+='\n元数据警告 '+str(len(result['warnings']))+' 条'
            peak=rec.get('peakWorkingSetBytes')
            rows.append([rec.get('name',''), str(rec.get('durationMs','—')), f'{peak/1048576:.2f}' if peak else '未捕获',actual[:80]])
        table(c,['操作','耗时 ms','峰值 MiB','实际结论'],rows,[6.0,2.3,2.5,5.2],9.5)
        strict_runs=[r for r in perf.get('records',[]) if isinstance(r.get('result'),dict) and r['result'].get('metadataStrict') is True]
        para(c,'数据集为重复的 1 MiB 模式块，块内含固定种子伪随机模式；存储比例同时受块去重与压缩影响，不能当作随机文件的普遍压缩率。'+('当前性能恢复为严格元数据模式，相关记录无警告。' if strict_runs and all(not x['result'].get('warnings') for x in strict_runs) else '元数据模式和警告以每条结果为准。')+'历史运行出现同值属主被重复设置导致 Win32 5 的警告，原记录保留于 evidence/benchmark-before-metadata-fix.json；修复后性能重跑另存，不能抹去历史异常。')
        recheck=load_json(ROOT/'evidence/metadata-recheck-da8c16f441144e5d9ac35136cb95777f/result.json')
        if recheck:
            para(c,f'修复后另作 D 盘严格元数据复验，恢复退出码 {recheck.get("restoreExit")}，警告 {recheck.get("warningCount")} 条。属主 SID、组 SID、实际 DACL ACE 及继承保护状态分别比较一致。Windows 会规范化 AUTO_INHERITED 来源标志，原始 SDDL 多出 AI，因此不声称安全描述符所有控制位逐位相同；复验 JSON 同时保留该差异。')
    else: para(c,'长数据集性能运行尚未产生最终汇总，不能以示例数字代替实际结果。已有超过 4 GiB 稀疏文件单元记录只证明大长度与稀疏语义，不替代 1 GiB 连续数据的吞吐测量。')
    net=load_json(ROOT/'evidence/network-smoke-results.json')
    records=net.get('records',[])
    comparison=next((x for x in records if x.get('name')=='network-transfer-comparison'),{})
    if comparison:
        table(c,['真实网络操作','对象负载 字节','含控制帧发送 字节'],[
            ['首次上传',comparison.get('initialPayloadBytes'),comparison.get('initialFramedBytes')],
            ['原样再次上传',comparison.get('repeatPayloadBytes'),comparison.get('repeatFramedBytes')],
            ['修改一个块后上传',comparison.get('incrementalPayloadBytes'),comparison.get('incrementalFramedBytes')]], [6.1,4.9,5.0])
        para(c,'重复上传的对象负载为零仍有握手和控制报文，报告把两种统计分开。局部修改还需传输新快照清单，因此实际负载略高于一个数据块。此处是真实本机 TCP 测量，跨机器局域网速度会受硬件与网络影响，不能直接等同。')
    heading(c,'5 健壮性和安全测试',2)
    times=load_json(ROOT/'evidence/metadata-times-results.json')
    if times:
        para(c,'三类时间另由 scripts/metadata-time-smoke.ps1 做独立 Win32 验收。给文件和空目录预置互不相同的创建、访问、修改时间，严格恢复后先读取原始 FILETIME 整数，再读取内容，避免验证动作改变访问时间。evidence/metadata-times-results.json 记录六个时间字段分别一致、严格模式及零警告；该独立验收不与 C++ 用例重复累计。')
    para(c,'恢复保护测试先在目标写入不同内容，再验证默认拒绝覆盖和显式替换；损坏块预检要求目标目录尚未创建。写入中途故障注入在第一块已经写入临时文件后抛错，检查旧文件与临时文件清理。真实权限用例配置拒绝读取的 DACL 并用受限令牌执行，避免管理员备份特权掩盖错误。')
    para(c,'网络测试直接向服务端发送错误密码、设备密钥、重放帧、篡改密文和超长长度，并验证普通账户管理拒绝。服务端授权不能只靠隐藏 GUI 按钮。纠错测试覆盖超过能力、错误侧车、头尾索引均坏及有效副本冲突，必须拒绝无法确认的修复结果。')
    issues=[x for x in cases if x['result']!='PASS']
    if issues:
        issues_cell=bound_table_heading(c,'6 结果分析和缺陷记录',2)
        table(issues_cell,['测试','状态','实际原因'],[(x['name'],x['result'],x.get('detail','')) for x in issues],[6.2,1.7,8.1],9.5)
    else:
        heading(c,'6 结果分析和缺陷记录',2)
        para(c,'当前自动测试未记录失败或跳过。该结论仅适用于所记录环境和输入，不等同于所有部署与权限场景均已验证。')
    para(c,'开发过程中已发现并处理的具体问题包括：账户名中的反斜杠被路径规则误转换；稳定硬链接在变化筛选中被重复判断为改变；最后恢复父目录 ACL 可能影响子对象；归档输出链接可能绕过源目标关系；设备名 CONIN$、CONOUT$ 和带上标数字的 COM 名称需额外拒绝。修复增加对应断言并通过真实运行记录验证，Git 留有独立提交。')
    para(c,'补充审查发现 PAX 扩展头填充没有验证零字节；先保留真实失败证据，再加入 TAR_PADDING 检查。Junction 捕获曾使用显示名而恢复使用 SubstituteName，现统一真实目标解析并检查原始重解析数据和清单一致性。调度补测修复损坏任务不可见、任务 ID 与文件名不一致、极端重试计数溢出及配置目录失效可能逃出完成回调的问题；故障隔离不允许坏任务阻止有效任务。通配筛选也改为完全自主匹配核心，并用 37 条实际文件用例回归预览和捕获一致性。')
    para(c,'带覆盖率插桩的 GUI 执行还暴露测试等待条件过早：文件已经原子提交，但工作线程尚未发出完成信号，测试即点击下一操作而遇到 busy 拒绝。修复仅调整测试工具为等待界面明确完成状态，未改生产功能；原 18 通过、2 失败记录保留于 evidence/history/coverage-gui-completion-race，修复后的结果单独记录，不删除失败历史。')
    para(c,'发布包后台进程复验首先发现测试观察窗口：Agent 原子替换 job.json 时，轮询测试短暂无法打开文件。测试工具仅对暂时文件打开/读取失败重试，仍直接暴露 JSON 损坏等其他异常。这一修复属于测试工具；失败历史与重跑分别归档。后续复跑进一步发现真实生产问题：Repository::list 没有持仓库锁，会与版本淘汰并发读取已删除的快照。修复给公开列举接口加锁，已持锁的 verify/prune 使用内部无重复加锁的辅助函数，避免死锁；新增回归区分这两类问题。')
    para(c,'现有边界包括：符号链接需要系统权限；非当前属主或限制性 ACL 可能需提升权限；已有目录 ADS 不直接覆盖；绝对链接保留原目标；变化筛选不是逐字节内容证明；网络暂存无自动过期清理；自研传输协议没有前向保密。宽松元数据模式只允许明确警告，不能据此声称完整恢复权限。')
    para(c,'恢复对静态父路径重解析点和每个对象创建前的祖先进行检查，但路径 API 不宣称能完整抵御另一本地恶意进程同时替换祖先目录。目标应由当前用户控制。普通文件逐个临时写入并验证后提交；整棵目录不具备全局回滚，后续失败时先前成功恢复的文件可能保留。')
    gui=load_json(ROOT/'evidence/gui-test-results.json'); gui_cases=gui.get('cases',[])
    if gui_cases:
        gui_cell=bound_table_heading(c,'7 GUI 和部署复验',2)
        table(gui_cell,['GUI 实际测试','结果','耗时 ms'],[(x['name'],x['result'],x.get('durationMs','')) for x in gui_cases],[11.3,2.2,2.5],9.5)
    else: heading(c,'7 GUI 和部署复验',2)
    background=load_json(ROOT/'evidence/background-test-results.json'); process_cases=background.get('cases',[])
    if process_cases:
        para(c,'跨进程复验启动真实 ChengVaultAgent 与 GUI，先保存任务，再关闭 GUI 后修改源文件。测试要求独立 Agent 持续运行、产生新版本、严格恢复新增内容，随后淘汰只保留两个可验证版本；同时检查重复后台进程被互斥锁拒绝。以下结果来自 evidence/background-test-results.json。')
        table(c,['真实后台进程验证','结果','耗时 ms'],[(x['name'],x['result'],x.get('durationMs','')) for x in process_cases],[11.3,2.2,2.5],9.5)
    deployment=load_json(ROOT/'evidence/deployment-results.json')
    if deployment:
        para(c,'发布部署验证将子进程 PATH 限为 '+deployment.get('childPath','')+'，清除开发环境 Qt 插件配置，再启动发布目录中的程序。此项为当前机器隔离运行环境复验，independentMachine=false，不能替代独立 Windows 机器。重复执行同一 GUI/后台用例用于部署验证，不计为新增用例扩大总数。')
        table(c,['发布包真实操作','结果','独立证据'],[(x['name'],'PASS' if x.get('passed') else 'FAIL',str(x.get('cases',''))+' 用例' if 'cases' in x else str(x.get('durationMs','—'))+' ms') for x in deployment.get('records',[])],[8.0,2.0,6.0],9.5)
    picture(c,ROOT/'evidence/gui/network.png','图 11 网络仓库操作界面')
    picture(c,ROOT/'evidence/gui/automation.png','图 12 自动任务界面')
    demo=load_json(ROOT/'evidence/demo/latest.json')
    demo_run=Path(demo.get('runDirectory','')).name or '未记录'
    para(c,'GUI 图片来自实际程序窗口，用于确认导航、输入区和操作布局，不单独证明所有交互已经通过自动化。最新真实演示 '+demo_run+'，时长 '+str(demo.get('durationSeconds','未记录'))+' 秒；记录包含损坏归档、纠错修复、密码恢复及内容校验。视频摘要及实际程序摘要由 evidence/demo/latest.json 绑定。独立 Windows 机器的复验与权限差异仍需另附记录，没有执行的项目不得填为通过。')
    uml=load_json(ROOT/'evidence/uml-native/validation.json')
    para(c,'StarUML 7.1.1 原生程序已安装，模型另通过官方元模型静态检查。原生加载与官方 API 渲染结果以 evidence/uml-native/validation.json 为准：'+('原生打开及七张图渲染已验证；未注册版导出带厂商水印，不能宣称这些原生 PNG 为出版级排版。正式报告使用同模型的原创清晰 SVG/PNG 插图，水印原生输出原样保留作为证据。' if uml.get('nativeStarUmlOpenVerified') else '原生打开尚未完成验证，不把静态检查写成原生软件验收。'))

    c=cell(13); heading(c,'七 AI 开发工具使用总结',1)
    para(c,'本项目使用 Codex 辅助需求整理、接口协商、C++ 模块实现、测试构造、编译诊断、文档排版及证据汇总。主要方式是先固定需求与模块边界，再分派密码/纠错、文件系统/归档、仓库/网络及界面/调度任务，通过共享头文件协作，并以编译、标准向量和真实系统测试核查输出。')
    table(c,['使用环节','作用','人工必须完成的复核'],[
    ['需求与设计','整理课程约束、追踪表、UML 与接口','确认范围和平台边界，能解释设计取舍'],
    ['算法和代码','自主实现与多轮编译修正，补充边界处理','理解算法，不把生成代码的数量当作掌握程度'],
    ['测试与诊断','标准向量、NTFS 和真实 TCP 自动测试','审查输入和期望，独立机器及权限复验'],
    ['报告和答辩','从实际 JSON 生成结果、组织图表和流程','补真实成员经历、操作截图，现场独立讲解']], [3.0,6.3,6.7])
    para(c,'效果评价：模块并行降低重复实现，标准向量与真实 API 测试能及时发现功能和边界问题。')
    para(c,'AI 使用操作截图、每位成员实际采用的提示和人工修改说明待小组按真实过程补充。推荐截图展示需求分解、一个失败测试及修复前后、人工复验，而不是仅展示最终生成内容。')
    c=cell(14); heading(c,'八 实验结论',1)
    para(c,'澄匣备份实现了 Windows NTFS 数据备份与恢复，并以独立模块提供筛选、两种打包、两种压缩、两种密码算法、自动任务、网络仓库和冗余纠错。关键架构是统一文件清单、块接口和完整快照；因此 GUI、CLI、后台与网络可以复用同一套恢复规则。')
    all_counts=counts+Counter(x['result'] for x in gui_cases)+Counter(x['result'] for x in process_cases)
    para(c,f'本次报告读取核心 {len(cases)} 个自动测试记录，其中 {counts["PASS"]} 通过、{counts["FAIL"]} 失败、{counts["SKIP"]} 跳过；加入 GUI 和真实后台进程后合计 {all_counts["PASS"]} 通过、{all_counts["FAIL"]} 失败、{all_counts["SKIP"]} 跳过。结果支持已执行环境下的功能与算法结论；未通过或未执行事项必须按测试章节继续处理。课程分数、两种归档的分别计分以及纠错额外加分以教师认定为准，本报告不承诺得分。')
    c=cell(15); heading(c,'九 总结及心得体会',1)
    para(c,'工程复盘表明，备份软件的难点在于失败条件下仍然保持可信。文件链接、稀疏洞和命名流不能仅靠复制主流处理；密码算法必须使用独立已知答案；恢复要把目标检查放在实际写入之前。统一清单与块接口使这些规则可以被归档和版本仓库共同使用，减少了修复一处却遗漏另一入口的风险。')
    para(c,'协作上的有效做法是先约定接口和模块所有权，再分模块提交；跨模块变更通过实际编译和集成测试确认。开发记录保留失败与修复，不把“功能计划已经写出”混同于程序已通过。文档也从真实 JSON 生成，避免测试数量和结果在多份材料中不一致。')
    para(c,'组长个人体会 待本人根据实际贡献填写。组员一个人体会 待本人填写。组员二个人体会 待本人填写。应说明各自负责内容、遇到的问题、具体修改与验证结果，不以统一生成的个人感受替代真实学习过程。')
    c=cell(16); heading(c,'十 实验过程与方法的改进建议',1)
    para(c,'后续首先补充独立 Windows 环境、符号链接权限和更复杂 ACL 的复验，完善覆盖率低分支的针对性测试。性能方面可比较内容定义分块与固定分块的代价，并减少 PAX 稀疏归档的临时磁盘放大。安全方面可在不影响本次自主算法考核的前提下，单独研究成熟传输协议、密钥轮换与前向保密，不把课程协议直接用于高价值生产数据。')
    para(c,'课程推进可以把需求冻结、可运行基础版本、扩展验证、交付验收设为明确里程碑，并以 lg.team 和真实 Git 历史辅助复盘。评分证据宜逐项对应功能演示、代码位置和测试结果，特别区分已实现、已测试、权限受限与仍需人工补充的事项。')
    heading(c,'参考规范与项目证据',2)
    refs=[
    '课程材料 0.pdf、老师发的东西.txt、其他评分细则或者注意事项.txt、实验报告模板。',
    'NIST FIPS 197 Advanced Encryption Standard，https://csrc.nist.gov/pubs/fips/197/final。',
    'NIST SP800 38A AES CTR 测试向量；FIPS 180 4 SHA256 标准。',
    'RFC 8439 ChaCha20；RFC 4231 HMAC；RFC 5869 HKDF；RFC 7914 中 PBKDF2 SHA256 测试向量。',
    'RFC 5510 Reed Solomon Forward Error Correction，https://www.rfc-editor.org/rfc/rfc5510.html。',
    'Microsoft Learn Win32 文件流、重解析点、稀疏文件、SetSecurityInfo 与 ReadDirectoryChangesW 文档。',
    '项目 docs/manifest.md、archive-format.md、network.md、ntfs-support.md；源码与测试均为本项目产物。',
    'evidence/test-results.json、network-smoke-results.json、benchmark 目录、GUI截图及 docs/traceability.csv。']
    for i,ref in enumerate(refs,1): para(c,f'[{i}] '+ref,size=10,indent=False)
    # Word's compatibility mode ignores keepNext within this inherited long table
    # and may push whole rows onto blank pages. Join consecutive heading lines
    # with their following prose into one keepLines paragraph instead. Runs retain
    # their own sizes and weight; table geometry and the original cover stay intact.
    for idx in range(4,17):
        tc=cell(idx)._tc; children=list(tc); j=0
        while j<len(children):
            first=children[j]
            outline=first.xpath('./w:pPr/w:outlineLvl') if first.tag==qn('w:p') else []
            if first.tag!=qn('w:p') or not outline or outline[0].get(qn('w:val'))=='0' or first.xpath('.//w:drawing'):
                j+=1; continue
            group=[first]; k=j+1
            while k<len(children) and children[k].tag==qn('w:p') and children[k].xpath('./w:pPr/w:outlineLvl') and not children[k].xpath('.//w:drawing'):
                group.append(children[k]); k+=1
            if k<len(children) and children[k].tag==qn('w:p') and not children[k].xpath('.//w:drawing'):
                following=children[k]
                if ''.join(following.itertext()).strip(): group.append(following); k+=1
            if len(group)>1:
                for part in group[1:]:
                    run=OxmlElement('w:r'); run.append(OxmlElement('w:br')); first.append(run)
                    for node in list(part):
                        if node.tag!=qn('w:pPr'): first.append(node)
                    tc.remove(part)
                # Mixed title/body blocks read from the left; individual runs
                # preserve the original title weight and hierarchy.
                from docx.text.paragraph import Paragraph
                p=Paragraph(first,cell(idx)); p.alignment=WD_ALIGN_PARAGRAPH.LEFT
                p.paragraph_format.keep_together=True
            j=k
    # All source cell edits are complete. Reuse existing XML package parts and
    # preserve every opaque part that python-docx could otherwise normalize.
    raw=work/'generated.docx'; doc.save(str(raw))
    original=zipfile.ZipFile(REFERENCE); generated=zipfile.ZipFile(raw)
    editable={'word/document.xml','word/_rels/document.xml.rels','[Content_Types].xml'}
    final=OUTPUT/'澄匣备份ChengVault实验报告.docx'
    with zipfile.ZipFile(final,'w',zipfile.ZIP_DEFLATED) as out:
        original_names=set(original.namelist())
        for name in original.namelist(): out.writestr(name,generated.read(name) if name in editable else original.read(name))
        for name in generated.namelist():
            if name not in original_names: out.writestr(name,generated.read(name))
    inventory=load_json(ROOT/'.cache/template-render/package-inventory.json',[])
    preserved=[]
    with zipfile.ZipFile(final) as output:
        for item in inventory:
            if item['classification']=='preserve':
                digest=hashlib.sha256(output.read(item['path'])).hexdigest()
                assert digest==item['sha256'], 'Preserve-only part modified: '+item['path']
                preserved.append(item['path'])
    assert hashlib.sha256(REFERENCE.read_bytes()).hexdigest()==EXPECTED_REFERENCE
    result={'output':str(final),'generatedChinaTime':now,'testEvidence':str(ROOT/'evidence/test-results.json'),
            'testCounts':dict(counts),'detailedCases':len(covered),'allCases':len(cases),'preservedParts':preserved,
            'referenceSha256':EXPECTED_REFERENCE,'templateOuterRows':18,'referenceSectionCount':len(doc.sections)}
    (work/'build-result.json').write_text(json.dumps(result,ensure_ascii=False,indent=2),encoding='utf-8')
    print(json.dumps(result,ensure_ascii=False,indent=2))

if __name__=='__main__': build()
