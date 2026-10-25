from pathlib import Path
import csv
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(r"D:\Codes\soft-project\ChengVault-team")
OUTPUT = Path(r"D:\Codes\soft-project\报告文档\ChengVault三人分工与提交指南.docx")
HEADERS = ["任务编号","负责人","复核人","阶段","具体工作","最终代码与测试","文档与证据","验收条件","状态"]
TASKS = [
["A01","刘子健","许珂瑞","模块实现","文件采集、归档打包、文件测试、界面入口、需求文档","src/platform/filesystem.cpp;include/chengvault/filesystem.hpp;src/core/archive.cpp;include/chengvault/archive.hpp;src/gui/*;src/apps/gui_main.cpp;src/apps/cli_main.cpp;tests/filesystem_tests.cpp;tests/filesystem_adversarial_tests.cpp;tests/gui_smoke.cpp","docs/requirements.md;docs/manifest.md;docs/ntfs-support.md;docs/archive-format.md;docs/configuration.md;docs/user-manual.md;docs/uml/**","文件系统、归档和 UI 用例通过；8 页手工复验完成","已纳入最终提交"],
["B01","许珂瑞","符若晨","模块实现","公共工具、压缩模块、加密模块、纠错编码、算法文档","src/core/common.cpp;src/core/compression.cpp;src/core/crypto.cpp;src/core/repair.cpp;include/chengvault/common.hpp;compression.hpp;crypto.hpp;repair.hpp;tests/algorithms_tests.cpp;tests/test_main.cpp;tests/test_support.hpp","docs/algorithms.md;docs/repair-format.md;docs/quality-and-coverage.md;docs/traceability.csv","标准向量、压缩往返和 RS 组合用例通过；覆盖率证据一致","已纳入最终提交"],
["C01","符若晨","刘子健","模块实现","构建骨架、版本仓库、调度后台、网络仓库、脚本交付","CMakeLists.txt;CMakePresets.json;src/core/repository.cpp;src/core/scheduler.cpp;src/core/network.cpp;src/apps/agent_main.cpp;src/apps/server_main.cpp;include/chengvault/repository.hpp;scheduler.hpp;network.hpp;tests/scheduler_tests.cpp;tests/background_process.cpp;tests/network_repository_tests.cpp;tests/network_repository_cases.hpp","scripts/**;docs/delivery-checklist.md;docs/external-acceptance.md;docs/network.md;PPT 与演示视频","构建、测试、打包和真实双机局域网验收通过","已纳入最终提交"],
["A02","刘子健","许珂瑞","验收","界面复验、符号链接、当前机器独立复验","scripts/validate-ui.ps1;scripts/symlink-acceptance.ps1;scripts/deployment-smoke.ps1","evidence/ui-refresh/**;evidence/symlink-acceptance-results.json;evidence/independent-machine/**","8 页截图、符号链接 PASS、四程序启动和备份恢复通过","已纳入最终提交"],
["B02","许珂瑞","符若晨","验收","测试、覆盖率、质量与演示证据","scripts/test.ps1;scripts/coverage.ps1;scripts/quality.ps1","evidence/test-results.json;evidence/gui-test-results.json;evidence/background-test-results.json;evidence/coverage/**;evidence/quality/**;evidence/demo/latest.json","核心 591、GUI 20、后台 5，合计 616 通过；行覆盖 97.6%","已纳入最终提交"],
["C02","符若晨","许珂瑞","验收","真实双机局域网","scripts/external-acceptance.ps1","evidence/lan/20261025-two-machine/**;docs/external-acceptance.md","PASS_TWO_MACHINE_LAN；两台机器标识不同；下载恢复 SHA-256 一致","已纳入最终提交"],
["A03","刘子健","许珂瑞","报告","报告定稿与需求、设计、测试三章","报告文档/澄匣备份ChengVault实验报告.docx","docs/requirements.md;docs/uml/**;evidence/report/**","报告数字与最终 evidence 一致；目录、页码和书签齐全","已纳入最终提交"],
["B03","许珂瑞","符若晨","文档","算法文档与答辩问答","docs/algorithms.md;docs/repair-format.md;docs/quality-and-coverage.md","docs/traceability.csv;evidence/coverage/**","能解释 Huffman、LZ77、AES-256、ChaCha20、HMAC、KDF 和 RS 边界","已纳入最终提交"],
["C03","符若晨","刘子健","交付","PPT、讲稿、视频、最终打包","scripts/build_presentation.mjs;scripts/record-demo.ps1;scripts/package.ps1;scripts/verify-package.ps1","PPT/**;演示视频/**;dist/**;交付说明.txt;交付文件校验清单.json;ALL/**","最终包从真实提交构建；18 文件哈希匹配；tag v1.0-submission","已纳入最终提交"],
["ALL01","刘子健/许珂瑞/符若晨","三人交叉复核","集成","最终集成、测试和证据冻结","final/freeze;master","evidence/**;tag v1.0-submission","工作区干净；616 测试通过；证据与源码一致","已纳入最终提交"],
["ALL02","刘子健/许珂瑞/符若晨","三人交叉复核","提交","身份、AI 截图、lg.team、视频链接和 ALL 归档","报告文档/**;PPT/**;演示视频/**","lg.team 截图;AI 使用截图;提交回执","真实材料补齐；链接可访问；ALL 包完整","已纳入最终提交"],
]

def set_run(run, size=10.5, bold=False):
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = RGBColor(0,0,0)
    rPr = run._element.get_or_add_rPr(); rf = rPr.get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), "宋体")
    for a in ("w:ascii","w:hAnsi","w:cs"): rf.set(qn(a), "Times New Roman")

def para(doc, text, size=10.5, bold=False, after=5):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(after); p.paragraph_format.line_spacing=1.25
    set_run(p.add_run(text), size, bold); return p

def heading(doc, text, level=1):
    p=doc.add_paragraph(); p.paragraph_format.space_before=Pt(12 if level==1 else 8); p.paragraph_format.space_after=Pt(5); p.paragraph_format.keep_with_next=True
    set_run(p.add_run(text), 14 if level==1 else 12, True); return p

def shade(cell):
    tcPr=cell._tc.get_or_add_tcPr(); shd=OxmlElement("w:shd"); shd.set(qn("w:fill"),"E7EFEA"); tcPr.append(shd)

def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers)); t.style="Table Grid"; t.alignment=WD_TABLE_ALIGNMENT.CENTER; t.autofit=False
    for i,h in enumerate(headers):
        c=t.rows[0].cells[i]; c.width=Cm(widths[i]); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER; shade(c)
        p=c.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_run(p.add_run(h),10.5,True)
    for row in rows:
        cells=t.add_row().cells
        for i,value in enumerate(row):
            cells[i].width=Cm(widths[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cells[i].paragraphs[0]; set_run(p.add_run(str(value)),10.5,False)
    return t

def footer_page(doc):
    footer=doc.sections[0].footer; p=footer.paragraphs[0]; p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    set_run(p.add_run("第 "),10.5)
    begin=OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"),"begin")
    instr=OxmlElement("w:instrText"); instr.set(qn("xml:space"),"preserve"); instr.text=" PAGE "
    sep=OxmlElement("w:fldChar"); sep.set(qn("w:fldCharType"),"separate")
    val=OxmlElement("w:t"); val.text="1"
    end=OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"),"end")
    r=p.add_run(); r._r.append(begin); r._r.append(instr); r._r.append(sep); r._r.append(val); r._r.append(end)
    set_run(p.add_run(" 页"),10.5)

doc=Document()
sec=doc.sections[0]; sec.page_width=Cm(21); sec.page_height=Cm(29.7); sec.left_margin=Cm(2.2); sec.right_margin=Cm(2.2); sec.top_margin=Cm(2.2); sec.bottom_margin=Cm(2.2)
style=doc.styles["Normal"]; style.font.name="Times New Roman"; style.font.size=Pt(10.5)
style._element.rPr.rFonts.set(qn("w:eastAsia"),"宋体")
footer_page(doc)

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; set_run(p.add_run("ChengVault 三人分工与提交指南"),18,True)
heading(doc,"一、最终完成情况")
para(doc,"GitHub 仓库：https://github.com/Ey-in/ChengVault.git，公开仓库。三人按 A/B/C 分支提交，经交叉评审后合并到 master。")
para(doc,"最终测试：核心 591 通过、GUI 20 通过、后台 5 通过，合计 616 通过，0 失败，0 跳过。覆盖率：行 3694/3786=97.6%，原始 GCC 分支 8184/14446=56.7%，辅助分支 86.5%，判定覆盖 671/888=75.6%。")
para(doc,"符号链接验收 PASS 2/0/0；机器独立复验已记录；两台真实 Windows 机器局域网验收为 PASS_TWO_MACHINE_LAN。演示视频为 2 分 58 秒、1280×720。")
para(doc,"最终交付包括最终报告、PPT、演示视频、运行包、源码包、history bundle、18 文件校验清单和 ALL 归档。")

heading(doc,"二、三人模块分工")
table(doc,["成员","最终代码与测试","文档与交付"],[
["A 刘子健\n组长","src/platform/filesystem.cpp；src/core/archive.cpp；src/gui/*；src/apps/gui_main.cpp、cli_main.cpp；tests/filesystem_tests.cpp、filesystem_adversarial_tests.cpp、gui_smoke.cpp","docs/requirements.md、manifest.md、ntfs-support.md、archive-format.md、configuration.md、user-manual.md；docs/uml/**；报告需求与设计章节"],
["B 许珂瑞","src/core/common.cpp、compression.cpp、crypto.cpp、repair.cpp；tests/algorithms_tests.cpp、test_main.cpp、test_support.hpp","docs/algorithms.md、repair-format.md、quality-and-coverage.md、traceability.csv；覆盖率、质量和演示证据"],
["C 符若晨","CMakeLists.txt、CMakePresets.json；src/core/repository.cpp、scheduler.cpp、network.cpp；src/apps/agent_main.cpp、server_main.cpp；tests/scheduler_tests.cpp、background_process.cpp、network_repository_tests.cpp、network_repository_cases.hpp","scripts/**；docs/delivery-checklist.md、external-acceptance.md、network.md；PPT、视频、最终打包和 18 文件清单"],
],[3.0,8.3,4.9])
para(doc,"复核关系：A 的提交由 B 复核和合并；B 的提交由 C 复核和合并；C 的提交由 A 复核和合并。跨模块接口由双方会审，不能各自改同一契约后直接合并。",10.5,False,8)

heading(doc,"三、提交与评审")
para(doc,"每名成员从 master 拉自己的分支，只提交本人负责的模块。分支示例：feature/member-a/gui、feature/member-b/docs、feature/member-c/network。")
para(doc,"提交流程：")
for line in ["git switch master","git pull --ff-only","git switch -c feature/member-a/gui","git status --short","git add -- <本人文件>","git diff --cached","git commit -m '<模块范围与验证结果>'","git push -u origin feature/member-a/gui"]:
    para(doc,line,10.5,False,1)
para(doc,"提交前不加入 build-*、dist、.toolchain、演示视频、connection-private.json、server-session.json、*.password 或 *.key。",10.5,False,8)

heading(doc,"四、验收与证据")
table(doc,["事项","负责人","验收条件","证据"],[
["界面复验","A 主责，B 复核","8 页、筛选、恢复、错密码、冲突、取消、后台淘汰和修复；截图与实际程序一致","evidence/ui-refresh/**"],
["符号链接与严格元数据","A 主责，B 复核","PASS 2/0/0；链接属性和目标属性一致；历史 SKIP 保留","evidence/symlink-acceptance-results.json"],
["当前机器独立复验","A 主责，C 复核","四程序启动；备份恢复哈希一致；记录机器标识与程序摘要","evidence/independent-machine/**"],
["真实双机局域网","C 主责，B 复核","PASS_TWO_MACHINE_LAN；双账户隔离、首次/重复/增量上传、下载恢复","evidence/lan/20261025-two-machine/**"],
["完整测试","B 主责，A 复核","616 通过、0 失败、0 跳过；测试名称与 evidence 一致","evidence/test-results.json 等"],
["覆盖率与质量","B 主责，C 复核","行 97.6%、原始分支 56.7%、辅助 86.5%、判定 75.6%；不设自设合格线","evidence/coverage/**"],
["最终打包","C 主责，A 复核","运行包、源码包、history bundle、package-verification.json、18 文件清单","dist/**；交付文件校验清单.json；ALL/**"],
["报告与 PPT","A 主责，B/C 复核","报告 QA、目录页码书签、PPT 20 页、逐页讲稿、2 分 58 秒视频","报告文档/**；PPT/**；演示视频/**"],
],[3.6,2.8,7.0,2.8])

heading(doc,"五、答辩分工")
table(doc,["成员","讲解内容","必须能回答"],[
["A 刘子健","需求、界面、筛选、NTFS 类型与元数据、索引归档和 PAX/TAR","筛选如何进入 capture；恢复顺序；为什么目录权限最后设置；路径逃逸和覆盖保护"],
["B 许珂瑞","Huffman、LZ77、AES-256-CTR、ChaCha20、HMAC、KDF、RS(12,10)","认证为何先于解密；标准向量为何必要；RS 每组最多两个已定位坏片；原始分支覆盖为何是 56.7%"],
["C 符若晨","版本仓库、调度后台、网络认证、增量与续传、构建和交付","GUI 退出后谁执行任务；重放如何拒绝；重复块为何不再上传；列表和淘汰为何共用仓库锁"],
],[3.0,7.2,6.0])

heading(doc,"六、提交清单")
table(doc,["项目","要求"],[
["报告","DOCX/PDF 与最终代码、evidence 一致"],
["答辩材料","PPT、逐页讲稿、演示流程"],
["源码与程序","源码包、运行包、history bundle；运行包从真实提交构建，checksums 与 package-verification 一致"],
["证据","test-results、gui/background、coverage、quality、deployment、ui-refresh、lan、symlink、independent-machine、demo/latest"],
["过程材料","lg.team 看板截图、PR 与评审记录、AI 使用截图"],
["提交包","交付说明.txt、交付文件校验清单.json；18 个文件哈希全部匹配；ALL1/ALL2 完整"],
],[4.0,12.2])

heading(doc,"七、一致性要求")
para(doc,"报告、PPT、指南和清单统一使用最终数字：616 通过、0 失败、0 跳过；行覆盖 97.6%，原始分支 56.7%，辅助分支 86.5%，判定覆盖 75.6%。")
para(doc,"私有文件只留在本机：connection-private.json、server-session.json、账户密码、设备密钥、构建目录、toolchain 和原始大体积测试数据。")

# Write CSV in UTF-8.
csv_path=ROOT/"docs"/"team-work-plan.csv"
with csv_path.open("w",encoding="utf-8-sig",newline="") as f:
    writer=csv.writer(f); writer.writerow(HEADERS); writer.writerows(TASKS)

OUTPUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUTPUT)
print("guide saved", OUTPUT)
print("csv saved", csv_path)